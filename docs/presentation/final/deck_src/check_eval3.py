"""SalesLuv_03_검증과확장.pptx 검수 — 수치 출처 · 페이지 번호 · 레이아웃.

각 장에 적힌 숫자가 그 장의 근거 평가 pptx 에 실제로 있는지 대조한다.
슬라이드마다 근거 파일을 하나로 못박아, 다른 평가의 숫자가 섞이면 실패한다.
"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from build_eval3 import (
    AT, DATA_X, DATA_W, FLOW_W, KPI_H, KPI_Y, LEAD_Y, OUT, ROW_H, ROW_Y,
)
from textwidth import width
from base_v3 import BOT, CW, FOOT_Y, M, SW, TOP

EVAL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "eval")
SOURCES = {1: "SalesLuv_미팅브리핑_평가.pptx",
           2: "SalesLuv_보고서작성에이전트_평가_4p.pptx",
           3: "SalesLuv_딜승산예측모델_평가.pptx"}
PAGES = {1: AT + 1, 2: AT + 2, 3: AT + 3}
NUM = re.compile(r"\d[\d,.]*")
# 근거 덱에 숫자로 적혀 있지 않고 본문 서술에서 끌어온 값. 확인한 근거를 함께 적는다.
DERIVED = {1: {"1": "1 · 3 · 5 · 10 · 30회 — 1p 평가 축",
               "3": "1 · 3 · 5 · 10 · 30회 — 1p 평가 축",
               "5": "1~5회 시점 평균 · 1p 평가 축"},
           2: {"12": "1p '일일 12 · 주간 4 · 월간 1'",
               "4": "1p '일일 12 · 주간 4 · 월간 1'",
               "1": "1p '일일 12 · 주간 4 · 월간 1'"},
           3: {"4": "2p '행마다 Unknown 을 4개로 맞춘다'",
               "10": "2p '448건 × 10세트'",
               "7": "5p 후보 7종 (RF 튜닝 … Stacking_LR)",
               "2": "5p 'Brier 1위 · FP 2위'",
               "1": "5p 'Brier 1위 · FP 2위'"}}


def texts(path):
    """평가 덱의 모든 텍스트 + 차트 값을 한 덩어리로 모은다."""
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif getattr(sh, "has_chart", False):
                ch = sh.chart
                out.extend(str(c) for c in ch.plots[0].categories)
                for ser in ch.plots[0].series:
                    for pt in ser._element.iter():
                        if pt.tag.endswith("}v") and pt.text:
                            out.append(pt.text)
            elif sh.has_table:
                out.extend(c.text for r in sh.table.rows for c in r.cells)
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)
    for s in Presentation(path).slides:
        walk(s.shapes)
    return "\n".join(out)


def canon(tok):
    t = tok.replace(",", "")
    if "." in t:
        t = t.rstrip("0").rstrip(".")
    return t


def numbers(blob):
    got = set()
    for tok in NUM.findall(blob):
        got.add(canon(tok))
        if "." in tok:                       # 0.18797 ↔ 0.1880
            for nd in (3, 4):
                got.add(canon("%.*f" % (nd, float(tok.replace(",", "")))))
    return got


def slide_text(sl):
    return "\n".join(sh.text_frame.text for sh in sl.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip())


def boxes(sl):
    return [(sh, sh.left / 914400, sh.top / 914400,
             (sh.width or 0) / 914400, (sh.height or 0) / 914400) for sh in sl.shapes]


def main():
    prs = Presentation(OUT)
    slides = list(prs.slides)
    fail = []

    # ---------- 1. 슬라이드 수 · 화면 비율 ----------
    assert len(slides) == 3, "슬라이드 %d 장 (3 이어야 함)" % len(slides)
    assert (round(prs.slide_width / 914400, 3), round(prs.slide_height / 914400, 3)) == (13.333, 7.5)

    # ---------- 2. 페이지 번호 (v3 삽입 위치 기준 30 · 31 · 32) ----------
    for n, sl in enumerate(slides, 1):
        found = None
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.top is None:
                continue
            if abs(sh.left - Inches(SW - M - 1.0)) < Inches(0.05) and \
               abs(sh.top - Inches(FOOT_Y - 0.03)) < Inches(0.05):
                found = sh.text_frame.text.strip()
        if found != "%02d" % PAGES[n]:
            fail.append("%d 장 페이지 번호 = %r (%02d 이어야 함)" % (n, found, PAGES[n]))

    # ---------- 3. 고정 제목 ----------
    TITLES = {1: "미팅 브리핑 에이전트", 2: "보고서 작성 에이전트", 3: "딜 승산 예측 모델"}
    for n, sl in enumerate(slides, 1):
        head = [sh.text_frame.text.strip() for sh, x, y, w, h in boxes(sl)
                if sh.has_text_frame and abs(y - 0.96) < 0.02]
        if head != [TITLES[n]]:
            fail.append("%d 장 제목 = %r" % (n, head))

    # ---------- 4. 수치 대조 ----------
    for n, fname in SOURCES.items():
        allowed = numbers(texts(os.path.join(EVAL, fname)))
        body = slide_text(slides[n - 1])
        body = body.replace("출처: SalesLuv 보고서작성 에이전트 평가", "")   # 파일명 속 '4p'
        for tok in NUM.findall(body):
            c = canon(tok)
            if c in allowed or c in DERIVED.get(n, {}):
                continue
            if c == str(PAGES[n]):             # 페이지 번호
                continue
            if re.fullmatch(r"0\d", tok):      # 01~09 = 챕터 · 단계 장식 번호
                continue
            fail.append("%d 장의 %r 을 %s 에서 찾지 못함" % (n, tok, fname))

    # ---------- 5. 다른 평가의 숫자가 섞였는지 ----------
    keys = {1: ["89.8", "96.3", "180", "1.7만"],
            2: ["95.4", "10.01", "4.57", "76.3"],
            3: ["0.1880", "16.15", "4.01", "4,480"]}
    for n in keys:
        for other, vals in keys.items():
            if other == n:
                continue
            for v in vals:
                if v in slide_text(slides[n - 1]):
                    fail.append("%d 장에 %d 장의 값 %r 이 있다" % (n, other, v))

    # ---------- 6. 레이아웃 ----------
    for n, sl in enumerate(slides, 1):
        for sh, x, y, w, h in boxes(sl):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:      # 로고
                continue
            if y > BOT:                                      # 출처 · 페이지 번호
                continue
            if x < M - 0.01 or x + w > M + CW + 0.01:
                fail.append("%d 장 좌우 여백 벗어남: %.2f~%.2f" % (n, x, x + w))
            if y + h > BOT + 0.01 and sh.has_text_frame and sh.text_frame.text.strip():
                fail.append("%d 장 본문 하단 넘침: %.2f" % (n, y + h))

        # 한 줄로 끝나야 하는 글상자가 칸을 넘지 않는가 (넘으면 줄바꿈 → 카드 밖으로 샌다)
        for sh, x, y, w, h in boxes(sl):
            if not (sh.has_text_frame and sh.text_frame.text.strip()) or not sh.width:
                continue
            if h > 0.36 or y > BOT:            # 여러 줄을 허용한 칸 · 푸터는 제외
                continue
            runs = sh.text_frame.paragraphs[0].runs
            if not runs:
                continue
            pt = runs[0].font.size.pt if runs[0].font.size else 15
            wt = 700 if runs[0].font.bold else 400
            txt = sh.text_frame.text.strip()
            if width(txt, pt, wt) > w + 0.01:
                fail.append("%d 장 한 줄 넘침: %r (%.2f > %.2f in)"
                            % (n, txt[:26], width(txt, pt, wt), w))

        # 글자끼리 겹치지 않는가 (작은 쪽 면적의 30% 넘게 물리면 실패)
        tb = [(sh.text_frame.text.strip()[:18], x, y, w, h) for sh, x, y, w, h in boxes(sl)
              if sh.has_text_frame and sh.text_frame.text.strip() and sh.width]
        for i, (t1, x1, y1, w1, h1) in enumerate(tb):
            for t2, x2, y2, w2, h2 in tb[i + 1:]:
                ox = min(x1 + w1, x2 + w2) - max(x1, x2)
                oy = min(y1 + h1, y2 + h2) - max(y1, y2)
                if ox > 0 and oy > 0 and ox * oy > 0.3 * min(w1 * h1, w2 * h2):
                    fail.append("%d 장 글자 겹침: %r ↔ %r" % (n, t1, t2))

    # ---------- 7. 밴드 ----------
    for name, a, b in (("결론", LEAD_Y, LEAD_Y + 0.34), ("KPI", KPI_Y, KPI_Y + KPI_H),
                       ("본문", ROW_Y, ROW_Y + ROW_H)):
        if a < TOP - 0.01 or b > BOT + 0.01:
            fail.append("%s 밴드가 본문 영역 밖: %.2f~%.2f" % (name, a, b))
    if LEAD_Y + 0.34 > KPI_Y or KPI_Y + KPI_H > ROW_Y:
        fail.append("결론 · KPI · 본문 밴드가 겹친다")
    if M + FLOW_W > DATA_X or DATA_X + DATA_W > M + CW + 0.01:
        fail.append("좌우 2단이 겹치거나 넘친다")

    if fail:
        print("실패 %d 건" % len(fail))
        for f in fail:
            print("  -", f)
        sys.exit(1)
    print("통과 — 3장 · 제목 · 페이지 번호 · 수치 출처 · 레이아웃 이상 없음")
    for n, fname in SOURCES.items():
        print("  %02dp ← %s" % (PAGES[n], fname))


if __name__ == "__main__":
    main()
