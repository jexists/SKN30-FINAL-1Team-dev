"""SalesLuv_OCR_고도화_3장.pptx 검수 — 수치 출처 · 페이지 번호 · 레이아웃.

3장에 적힌 모든 숫자가 근거 PDF(OCR 고도화 평가 최종보고서)에 실제로 있는지 대조한다.
없는 숫자가 하나라도 있으면 실패한다 — "수치를 지어내지 않았다"의 기계적 증명이다.
레이아웃은 Pretendard 실측 폭으로 한 줄 넘침을, 상자 겹침으로 카드 충돌을 잡는다.

pptx 는 python3.12 에만, pypdf 는 backend/.venv 에만 있다. PDF 본문만 venv 로 뽑아 온다.
"""

import math
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from base_v3 import BOT, CW, FOOT_Y, H1_Y, M, SW, TOP
from build_ocr3 import AT, OUT
from textwidth import width

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, "..", "..", "..", ".."))
PDF = os.path.join(ROOT, "docs", "presentation", "eval", "OCR 고도화 평가 최종보고서.pdf")
VENV_PY = os.path.join(ROOT, "backend", ".venv", "bin", "python")

CHAPTER = "03 검증과 확장"
PAGES = {1: AT + 1, 2: AT + 2, 3: AT + 3}
TITLES = {1: "원본 PDF OCR만으로는 핵심 필드가 누락됐다",
          2: "누락된 문서만 PNG로 다시 읽고, 취약 필드는 남겼다",
          3: "핵심 구조화 문서 일치율 65.4% → 98.8%"}
NUM = re.compile(r"\d[\d,.]*")


def pdf_text(path):
    """근거 PDF 본문. pypdf 가 있는 venv 인터프리터에 위임한다."""
    code = ("import sys;from pypdf import PdfReader;"
            "sys.stdout.write('\\n'.join(p.extract_text() for p in PdfReader(sys.argv[1]).pages))")
    r = subprocess.run([VENV_PY, "-c", code, path], capture_output=True, text=True)
    if r.returncode != 0 or not r.stdout.strip():
        sys.exit("PDF 본문을 읽지 못했다: %s" % (r.stderr.strip() or "빈 결과"))
    return r.stdout


def canon(tok):
    t = tok.replace(",", "")
    if "." in t:
        t = t.rstrip("0").rstrip(".")
    return t


def numbers(blob):
    got = set()
    for tok in NUM.findall(blob):
        got.add(canon(tok))
        if "." in tok:
            for nd in (2, 3, 4):
                try:
                    got.add(canon("%.*f" % (nd, float(tok.replace(",", "")))))
                except ValueError:
                    pass
    return got


def slide_text(sl):
    return "\n".join(sh.text_frame.text for sh in sl.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip())


def boxes(sl):
    return [(sh, sh.left / 914400, sh.top / 914400,
             (sh.width or 0) / 914400, (sh.height or 0) / 914400) for sh in sl.shapes]


def main():
    prs = Presentation(OUT)
    slides = list(prs.slides)
    fail = []

    # ---------- 1. 슬라이드 수 · 화면 비율 ----------
    assert len(slides) == 3, "슬라이드 %d 장 (3 이어야 함)" % len(slides)
    assert (round(prs.slide_width / 914400, 3), round(prs.slide_height / 914400, 3)) == (13.333, 7.5)

    # ---------- 2. 페이지 번호 (v3 + Deep Agents 3장 뒤 60 · 61 · 62) ----------
    for n, sl in enumerate(slides, 1):
        found = None
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.top is None:
                continue
            if abs(sh.left - Inches(SW - M - 1.0)) < Inches(0.05) and \
               abs(sh.top - Inches(FOOT_Y - 0.03)) < Inches(0.05):
                found = sh.text_frame.text.strip()
        if found != "%02d" % PAGES[n]:
            fail.append("%d 장 페이지 번호 = %r (%02d 이어야 함)" % (n, found, PAGES[n]))

    # ---------- 3. 제목 위치 · 챕터 라벨 (v3 와 같은 좌표) ----------
    for n, sl in enumerate(slides, 1):
        head = [sh.text_frame.text.strip() for sh, x, y, w, h in boxes(sl)
                if sh.has_text_frame and abs(y - H1_Y) < 0.02]
        if head != [TITLES[n]]:
            fail.append("%d 장 제목 = %r" % (n, head))
        if CHAPTER not in slide_text(sl):
            fail.append("%d 장 챕터 라벨 없음" % n)

    # ---------- 4. 수치 대조 — 모든 숫자가 근거 PDF 에 있어야 한다 ----------
    allowed = numbers(pdf_text(PDF))
    for n, sl in enumerate(slides, 1):
        for tok in NUM.findall(slide_text(sl)):
            c = canon(tok)
            if c in allowed:
                continue
            if c == str(PAGES[n]):                 # 페이지 번호
                continue
            if re.fullmatch(r"0\d", tok):          # 03 = 챕터 장식 번호
                continue
            fail.append("%d 장의 %r 을 근거 PDF 에서 찾지 못함" % (n, tok))

    # ---------- 5. 레이아웃 ----------
    for n, sl in enumerate(slides, 1):
        for sh, x, y, w, h in boxes(sl):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:      # 로고
                continue
            if y > BOT:                                      # 출처 · 페이지 번호
                continue
            if x < M - 0.01 or x + w > M + CW + 0.01:
                fail.append("%d 장 좌우 여백 벗어남: %.2f~%.2f" % (n, x, x + w))
            if y < TOP - 0.01 and sh.has_text_frame and sh.text_frame.text.strip() \
               and abs(y - H1_Y) > 0.02 and y > 0.8:
                fail.append("%d 장 본문이 제목 영역 침범: %.2f" % (n, y))
            if y + h > BOT + 0.01 and sh.has_text_frame and sh.text_frame.text.strip():
                fail.append("%d 장 본문 하단 넘침: %.2f" % (n, y + h))

        # 줄바꿈까지 계산한 실제 글 높이가 칸을 넘지 않는가
        for sh, x, y, w, h in boxes(sl):
            if not (sh.has_text_frame and sh.text_frame.text.strip()) or not sh.width:
                continue
            if y > BOT:
                continue
            tf = sh.text_frame
            inner = w - (tf.margin_left + tf.margin_right) / 914400
            need = 0.0
            wrapped = []
            for para in tf.paragraphs:
                runs = para.runs
                if not runs:
                    continue
                txt = "".join(r.text for r in runs)
                f0 = runs[0].font
                pt = f0.size.pt if f0.size else 15
                wt = 700 if f0.bold else (600 if (f0.name or "").endswith("SemiBold") else
                                          (500 if (f0.name or "").endswith("Medium") else 400))
                nl = max(1, math.ceil(width(txt, pt, wt) / inner - 0.005))
                if nl > 1:
                    wrapped.append(txt[:20])
                need += (para.space_before.pt if para.space_before else 0) / 72
                need += nl * pt * (para.line_spacing or 1.3) / 72
            if need > h + 0.02:
                fail.append("%d 장 글 높이 넘침: %r (%.2f > %.2f in%s)"
                            % (n, sh.text_frame.text.strip()[:22], need, h,
                               " · 줄바꿈 " + " / ".join(wrapped) if wrapped else ""))

        # 글자끼리 겹치지 않는가
        tb = [(sh.text_frame.text.strip()[:18], x, y, w, h) for sh, x, y, w, h in boxes(sl)
              if sh.has_text_frame and sh.text_frame.text.strip() and sh.width]
        for i, (t1, x1, y1, w1, h1) in enumerate(tb):
            for t2, x2, y2, w2, h2 in tb[i + 1:]:
                ox = min(x1 + w1, x2 + w2) - max(x1, x2)
                oy = min(y1 + h1, y2 + h2) - max(y1, y2)
                if ox > 0 and oy > 0 and ox * oy > 0.3 * min(w1 * h1, w2 * h2):
                    fail.append("%d 장 글자 겹침: %r ↔ %r" % (n, t1, t2))

    if fail:
        print("실패 %d 건" % len(fail))
        for f in fail:
            print("  -", f)
        sys.exit(1)
    print("통과 — 3장 · 제목 · 페이지 번호 %d·%d·%d · 수치 출처(PDF) · 레이아웃 이상 없음"
          % (PAGES[1], PAGES[2], PAGES[3]))


if __name__ == "__main__":
    main()
