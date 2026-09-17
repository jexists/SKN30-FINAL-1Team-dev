"""SalesLuv 최종 발표 v3 — v2 디자인 시스템을 그대로 쓰되 제목 아래 회색 부제를 없앤 프레임.

v2(base_v2.py)와 달라진 점은 슬라이드 프레임 하나뿐이다.
  - 큰 제목 바로 아래의 회색 subtitle 을 없앤다
  - 제목과 본문 사이 간격을 다시 잡아 본문 영역을 2.04 ~ 6.62 로 넓힌다
색·폰트·카드·아이콘·페이지 번호 등 나머지 시각 언어는 v2 를 그대로 가져온다.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from base_v2 import *  # noqa: F401,F403  (v2 디자인 시스템 전체 재사용)
from base_v2 import (
    BLUE, BLUE_PALE, BORDER, CW, EYE_Y, GRAY, HAIRLINE, INK, L, LOGO, LOGO_H, LOGO_W,
    BLUE_600, C, LOGO_WHITE, M, MUTED, R, SUB, SW, T_EYE, T_H1, T_NOTE, WHITE, _bar, blank,
    card, line,
    para, write,
)

# ---------- v3 전용 지오메트리 ----------
H1_Y = 0.96
RULE_Y = 1.80
TOP = 2.04          # 본문 시작 (부제 없음)
BOT = 6.62
FOOT_Y = 6.86
BAND = BOT - TOP    # 4.58


def slide(prs, page, *, chapter=None, headline=None, source=None, dark=False):
    """제목 아래 회색 부제가 없는 v3 표준 본문 슬라이드. (slide, 본문 시작 y) 반환."""
    sl = blank(prs, dark=dark)
    fg = WHITE if dark else BLUE
    if chapter:
        _bar(sl, M, EYE_Y + 0.035, 0.085, 0.2, fg)
        write(sl, M + 0.22, EYE_Y, 7.0, 0.3,
              [(chapter, dict(size=T_EYE, weight=600, color=fg, line=1.25))])
    sl.shapes.add_picture(LOGO_WHITE if dark else LOGO,
                          Inches(SW - M - LOGO_W), Inches(EYE_Y - 0.02),
                          Inches(LOGO_W), Inches(LOGO_H))
    if headline:
        write(sl, M, H1_Y, CW, 0.7,
              [(headline, dict(size=T_H1, weight=700, color=(WHITE if dark else INK), line=1.2))])
    line(sl, M, RULE_Y, M + CW, RULE_Y, color=HAIRLINE)
    if source:
        write(sl, M, FOOT_Y, CW - 1.2, 0.3,
              [(source, dict(size=T_NOTE, weight=400, color=MUTED, line=1.35))])
    write(sl, SW - M - 1.0, FOOT_Y - 0.03, 1.0, 0.3,
          [("%02d" % page, dict(size=16, weight=600, color=(BLUE_PALE if not dark else SUB),
                                line=1.25))], align=R)
    return sl, TOP


def monogram(sl, x, y, size, letter):
    """공식 아이콘이 없는 기술용 자리 표시. 브랜드 마크가 아닌 중립 타일."""
    sh = card(sl, x, y, size, size, fill=GRAY, radius=0.06, shadow=False, line=BORDER)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, letter, size=15, weight=700, color=BLUE_600, line=1.0, align=C, first=True)
    return sh
