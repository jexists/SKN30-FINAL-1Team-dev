"""SalesLuv 최종 발표 덱 v3 — v2 디자인 유지 + 제목 아래 회색 부제 제거.

v2 대비 바뀐 것
  1. 모든 본문 슬라이드에서 제목 아래 회색 부제를 없애고 본문 영역을 다시 잡음
  2. 기술 스택 2장 → 1장 (공식 로고 중심, OCR 카테고리 추가)
  3. 시스템 아키텍처 2장 → 1장 (한 화면 다이어그램 + 배포 스트립)
  4. AI 기능(ML · OCR · Agent) 역할 구분 슬라이드 추가
  5. 일정 에이전트 현재 동작 도식 추가 (코드 확인 기준)
  6. OCR 평가를 'OCR 고도화 전후 최종 결과보고서(2026-09-17)' 수치로 교체
  7. 향후 계획을 4개 카드 Roadmap 한 장으로 확정
사실·수치는 deck_prompt.md 확정본과 첨부 OCR 보고서, 구현 코드만 사용한다.
"""

import os
import sys
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import cairosvg
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from base_v3 import (
    BAR, BLUE, BLUE_500, BLUE_600, BLUE_DEEP, BLUE_LIGHT, BLUE_PALE, BOT, BORDER, C, COL, CW,
    DARK, DARK_SUB, EYE_Y, GRAY, HAIRLINE, INK, L, LOGO, LOGO_H, LOGO_W, LOGO_WHITE, M, MUTED,
    R, SH, SUB, SW, T_BODY, T_CARD, T_KPI, T_NOTE, TINT, TINT2, TOP, WHITE, _bar, add_chart,
    arrow, blank, bullets, card, circle, divider, flat, ghost, grid, keyline, kpi, label, line,
    make_white_logo, monogram, new_deck, note, numcard, para, pill, placeholder, slide, table,
    write,
)
from pptx.dml.color import RGBColor

OUT = ("/Volumes/Jexists/skn30/SKN30-FINAL-1Team-dev/docs/presentation/final/"
       "SalesLuv_최종발표_260918_v3.pptx")
HERE = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.path.join(HERE, "icons")

CH1 = "01 문제 정의"
CH2 = "02 설계와 구현"
CH3 = "03 검증과 확장"
APX = "부록 · 검증 자료"
PH = "[입력 예정]"

BRAND = {"react": "149ECA", "typescript": "3178C6", "vite": "646CFF", "fastapi": "009688",
         "python": "3776AB", "sqlalchemy": "D71F00", "supabase": "1FA97A", "postgresql": "4169E1",
         "openai": "412991", "langchain": "1C3C3C", "scikitlearn": "F7931E", "docker": "2496ED",
         "githubactions": "2088FF", "amazonwebservices": "E8871A", "amazons3": "569A31",
         "amazonec2": "E8871A", "paddlepaddle": "0062B0"}
ICON_SRC = "https://cdn.jsdelivr.net/npm/simple-icons@13/icons/{}.svg"


def icon(slug):
    os.makedirs(ICON_DIR, exist_ok=True)
    path = os.path.join(ICON_DIR, slug + ".png")
    if not os.path.exists(path):
        req = urllib.request.Request(ICON_SRC.format(slug), headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=20) as r:
            svg = r.read().decode()
        svg = svg.replace("<svg ", '<svg fill="#%s" ' % BRAND.get(slug, "45515E"), 1)
        cairosvg.svg2png(bytestring=svg.encode(), write_to=path, output_width=320, output_height=320)
    return path


def put_icon(sl, slug, x, y, size):
    """정사각 원본을 정사각으로 넣어 비율이 깨지지 않게 한다."""
    return sl.shapes.add_picture(icon(slug), Inches(x), Inches(y), Inches(size), Inches(size))


def tech(sl, slug, name, x, y, w):
    """작은 공식 로고 + 기술명 한 줄. slug 가 None 이면 중립 모노그램 타일."""
    s = 0.34
    if slug:
        put_icon(sl, slug, x, y, s)
    else:
        monogram(sl, x, y, s, name[0].upper())
    write(sl, x + s + 0.16, y + 0.03, w - s - 0.2, 0.3,
          [(name, dict(size=17, weight=500, color=INK, line=1.25))])


def build():
    make_white_logo()
    prs = new_deck()
    pg = [0]

    def p():
        pg[0] += 1
        return pg[0]

    # ========================= 01 표지 =========================
    sl = blank(prs, dark=True)
    sl.shapes.add_picture(LOGO_WHITE, Inches(SW - M - LOGO_W), Inches(EYE_Y - 0.02),
                          Inches(LOGO_W), Inches(LOGO_H))
    write(sl, M, 1.95, 8.0, 0.34,
          [("SKN30 1Team · CassTerra", dict(size=17, weight=600, color=BLUE_LIGHT, line=1.25))])
    write(sl, M, 2.38, 10.0, 1.4, [("SalesLuv", dict(size=78, weight=700, color=WHITE, line=1.1))])
    write(sl, M, 3.85, 10.0, 0.5,
          [("영업의 흐름을 연결하는 멀티에이전트 CRM",
            dict(size=25, weight=500, color=DARK_SUB, line=1.35))])
    _bar(sl, M, 4.62, 2.0, 0.075, BLUE)
    write(sl, M, 4.92, 10.0, 1.1,
          [("CONNECT THE HISTORY.", dict(size=23, weight=700, color=WHITE, line=1.4)),
           ("APPROVE THE NEXT MOVE.", dict(size=23, weight=700, color=BLUE_LIGHT, line=1.4))])
    write(sl, M, 6.28, 10.0, 0.7,
          [("2026.09.18 (FRI) 15:00", dict(size=18, weight=600, color=WHITE, line=1.4)),
           ("천성배 · 정주애 · 박제섭 · 박지유",
            dict(size=18, weight=400, color=DARK_SUB, line=1.4, spacing_before=4))])
    _bar(sl, 0, SH - 0.16, SW, 0.16, BLUE)
    p()

    # ========================= 02 목차 =========================
    sl, top = slide(prs, p(), chapter="목차", headline="문제 → 구현 → 검증, 세 파트로 설명합니다")
    w3, xs = grid(3, 0.5)
    toc = [("01", "문제 정의", "왜 이 문제를 골랐는가",
            "팀 · 주제 · 문제\n시장 · 경쟁 · 해결책 · 기대효과"),
           ("02", "설계와 구현", "무엇을 어떻게 만들었는가",
            "시나리오 · 시연 · 기술 스택\n아키텍처 · 에이전트 · AI 기능"),
           ("03", "검증과 확장", "제대로 동작하는지 어떻게 확인했는가",
            "에이전트별 평가 · 일정 에이전트\n실패와 개선 · 향후 계획 · 결론")]
    for i, (n, t, lead, kw) in enumerate(toc):
        x = xs[i]
        if i:
            line(sl, x - 0.25, top + 0.1, x - 0.25, BOT - 0.4, color=HAIRLINE)
        ghost(sl, x, top, w3, n, size=66, color=BLUE_PALE)
        write(sl, x, top + 1.35, w3, 0.45, [(t, dict(size=27, weight=700, color=INK, line=1.25))])
        write(sl, x, top + 2.05, w3, 0.8,
              [(lead, dict(size=18, weight=600, color=BLUE_600, line=1.4))])
        line(sl, x, top + 2.95, x + w3 - 0.3, top + 2.95, color=HAIRLINE)
        lines = [(s, dict(size=18, weight=400, color=SUB, line=1.45,
                          spacing_before=(0 if j == 0 else 4))) for j, s in enumerate(kw.split("\n"))]
        write(sl, x, top + 3.18, w3 - 0.2, 1.2, lines)

    # ========================= 03 파트 디바이더 1 =========================
    divider(prs, "01", "문제 정의", "왜 이 프로젝트가 필요한가", p())

    # ========================= 04 팀 구성 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="네 명이 기획부터 배포까지 나눠 맡았다")
    w4, xs = grid(4, 0.32)
    team = [("천성배", ["P.M", "OCR", "자료요약 Agent"]),
            ("정주애", ["Full Stack", "Infra"]),
            ("박제섭", ["기술리더", "미팅분석 Agent", "업무보고 Agent"]),
            ("박지유", ["영업계약 Agent", "일정관리 Agent", "시연영상"])]
    for i, (name, roles) in enumerate(team):
        x = xs[i]
        placeholder(sl, x, top, w4, 2.3, "[PHOTO]", size=16)
        write(sl, x, top + 2.48, w4, 0.42, [(name, dict(size=23, weight=700, color=INK, line=1.25))])
        _bar(sl, x, top + 3.02, 0.5, 0.05, BLUE)
        for j, role in enumerate(roles):
            write(sl, x, top + 3.22 + j * 0.38, w4, 0.34,
                  [(role, dict(size=T_BODY, weight=(600 if j == 0 else 400),
                               color=(BLUE_600 if j == 0 else SUB), line=1.35))])

    # ========================= 05 프로젝트 주제 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="담당자의 고객 맥락을 팀이 이어받게 한다",
                    source="출처: 프로젝트 기획서(2026-08)")
    flat(sl, M, top, CW, 1.15, fill=TINT2)
    write(sl, M + 0.4, top + 0.32, CW - 0.8, 0.6,
          [("영업의 시작부터 끝까지, 한 번에 해결하는 멀티에이전트 CRM",
            dict(size=27, weight=700, color=BLUE_DEEP, line=1.3))], align=C)
    w3, xs = grid(3, 0.36)
    defs = [("국내 B2B 영업", "고객 · 딜 · 일정 · 보고가\n흩어져 있는 조직"),
            ("멀티에이전트", "판단 · 초안 · 제안을\n다섯 에이전트로 분담"),
            ("Human-in-the-loop", "AI는 초안까지,\n확정은 사람의 승인으로")]
    dy = top + 1.6
    for i, (t, d) in enumerate(defs):
        x = xs[i]
        _bar(sl, x, dy, w3, 0.05, BLUE if i == 2 else BLUE_PALE)
        write(sl, x, dy + 0.3, w3, 0.4, [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.5,
                          spacing_before=(0 if j == 0 else 2))) for j, s in enumerate(d.split("\n"))]
        write(sl, x, dy + 0.88, w3, 1.1, lines)
    note(sl, M, BOT - 0.4, CW,
         "CRM — 고객과 나눈 이야기를 기억하는 시스템        딜 — 하나의 거래 기회", size=16)

    # ========================= 06 문제 정의 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="영업 맥락은 다섯 지점에서 끊긴다",
                    source="출처: 프로젝트 기획서(2026-08) · 영업사원 출신 PM의 현장 경험")
    probs = [("01", "영업정보 분산", "고객 · 계약 · 매출이 따로 관리된다"),
             ("02", "영업흐름 단절", "리드 → 계약 → 매출이 이어지지 않는다"),
             ("03", "판단 공백", "진행 상황과 예측이 보이지 않는다"),
             ("04", "인수인계 맥락 소실", "관계와 노하우가 담당자와 함께 사라진다"),
             ("05", "범용 CRM 한계", "B2B 영업 현장에는 기능이 부족하다")]
    for i, (n, t, d) in enumerate(probs):
        y = top + i * 0.92
        act = i == 3
        flat(sl, M, y, CW, 0.78, fill=(TINT2 if act else TINT), accent=(BLUE if act else None))
        write(sl, M + 0.36, y + 0.17, 0.9, 0.46,
              [(n, dict(size=26, weight=700, color=(BLUE if act else BLUE_LIGHT), line=1.2))])
        write(sl, M + 1.4, y + 0.23, 3.2, 0.4,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        write(sl, M + 4.9, y + 0.25, CW - 5.25, 0.4,
              [(d, dict(size=T_BODY, weight=400, color=SUB, line=1.3))])

    # ========================= 07 문제 결론 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="결국 남는 문제는 암묵지다",
                    source="출처: 프로젝트 기획서(2026-08) · 영업사원 출신 PM의 현장 경험")
    flat(sl, M, top, CW, 1.7, fill=TINT2, accent=BLUE)
    write(sl, M + 0.45, top + 0.36, CW - 0.9, 1.2,
          [("고객 맥락과 영업 노하우는 담당자 개인의 암묵지로만 쌓인다",
            dict(size=26, weight=700, color=INK, line=1.3)),
           ("조직 자산으로 남지 않아 담당자가 바뀌면 함께 사라진다",
            dict(size=T_CARD, weight=400, color=SUB, line=1.35, spacing_before=10))])
    label(sl, M, top + 2.08, "이 판단의 근거")
    w3, xs = grid(3, 0.36)
    basis = [("프로젝트 기획서", "2026-08 작성\n문제 정의와 범위"),
             ("현장 경험", "영업사원 출신 PM의\n실무 경험에서 출발"),
             ("정량 근거는 없음", "설문 · 인터뷰 데이터 없이\n검증할 가설로 다룬다")]
    by = top + 2.56
    for i, (t, d) in enumerate(basis):
        flat(sl, xs[i], by, w3, 1.92, fill=(GRAY if i == 2 else TINT))
        write(sl, xs[i] + 0.3, by + 0.32, w3 - 0.6, 0.4,
              [(t, dict(size=T_CARD, weight=600, color=(MUTED if i == 2 else INK), line=1.25))])
        lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.45,
                          spacing_before=(0 if j == 0 else 2))) for j, s in enumerate(d.split("\n"))]
        write(sl, xs[i] + 0.3, by + 0.9, w3 - 0.6, 1.0, lines)

    # ========================= 08 시장 규모 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="국내 CRM 시장은 2033년까지 두 배 이상 커진다",
                    source="출처: Grand View Research(2026) · KISDI 2026(원자료 Gartner 2025) · 환율 2026-09-15 1,347.8원")
    cwid = 7.1
    label(sl, M, top, "국내 CRM 시장 규모 (조원)")
    add_chart(sl, COL, M - 0.2, top + 0.42, cwid, 3.55, ["2025년", "2026년", "2033년"],
              [("국내 CRM 시장", [2.51, 2.75, 5.46])], label_fmt="0.00", gap=110, y_max=6.0,
              axis_size=T_NOTE, label_size=16)
    rx = M + cwid + 0.5
    rw = CW - cwid - 0.5
    kpi(sl, rx, top, rw, "10.3%", "국내 CRM 연평균 성장률", "2026–2033", tint=True, h=1.85)
    kpi(sl, rx, top + 2.02, rw, "17.1%", "국내 SaaS 연평균 성장률",
        "2025–2029 · 2조 1,295억 → 4조 164억", color=BLUE_500, tint=True, h=1.95)
    note(sl, M, BOT - 0.35, CW,
         "CRM과 SaaS는 시장 정의가 달라 합산하지 않는다 · 영업 CRM 단독 시장 통계는 공개된 것이 없다", size=16)

    # ========================= 09 진입 조건 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="진입 조건은 근거 · 연결 · 승인 세 가지다",
                    source="출처: 각 사 공식 제품 페이지(2026-09-15 확인) · 팀 시장 조사")
    w3, xs = grid(3, 0.36)
    conds = [("01", "경쟁 현실", "AI 미팅 요약과 후속조치\n제안은 이미 존재한다"),
             ("02", "초기 ICP", "문서 · 견적 · 계약이\n복잡한 B2B 영업팀"),
             ("03", "차별화 조건", "근거 검색 + 데이터 연결\n+ 담당자 승인")]
    for i, (n, t, d) in enumerate(conds):
        numcard(sl, xs[i], top, w3, 3.3, n, t, d, active=(i == 2), num_size=44, title_lines=1)
    keyline(sl, top + 3.6, "성장하는 시장에서 우리가 고른 자리 — 근거를 대고, 데이터를 잇고, 사람이 승인한다")

    # ========================= 10 경쟁 비교 ① =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="기본 기능은 세 제품 모두 갖추고 있다",
                    source="출처: 각 사 공식 제품 페이지(2026-09-15 확인) · 기능군 판정은 팀 조사 기준")
    rows = [["제품", "고객 · 거래처 관리", "파이프라인 · 딜 관리", "영업 실적 · 매출 분석"],
            ["세일즈맵", "●", "●", "●"],
            ["핑거세일즈", "●", "●", "●"],
            ["세일즈인사이트", "●", "●", "△"],
            ["SalesLuv", "●", "●", "●"]]
    table(sl, M, top, CW, 3.9, rows, col_w=[3.3, 2.931, 2.931, 2.931], highlight_row=4)
    note(sl, M, top + 4.08, CW, "● 제공        △ 일부        – 공개 자료에서 확인되지 않음", size=16)

    # ========================= 11 경쟁 비교 ② =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="차이는 AI 유무가 아니라 승인이 순환하는 구조다",
                    source="출처: 각 사 공식 제품 페이지(2026-09-15 확인) · 기능군 판정은 팀 조사 기준")
    rows = [["제품", "AI 기능", "사용자 승인 후 업무 반영", "승인 결과 → 다음 행동 순환"],
            ["세일즈맵", "△", "–", "–"],
            ["핑거세일즈", "●", "●", "–"],
            ["세일즈인사이트", "–", "–", "–"],
            ["SalesLuv", "●", "●", "●"]]
    table(sl, M, top, CW, 3.9, rows, col_w=[3.3, 2.5, 3.146, 3.147], highlight_row=4)
    note(sl, M, top + 4.08, CW, "● 제공        △ 일부        – 공개 자료에서 확인되지 않음", size=16)

    # ========================= 12 차별점 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="SalesLuv가 더하는 것은 세 가지다",
                    source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    w3, xs = grid(3, 0.36)
    diffs = [("근거 기반 추천", "보고서와 자료를 검색해\n출처와 함께 제안한다", BLUE_PALE),
             ("Human-in-the-loop", "AI 제안은 초안일 뿐,\n담당자가 확정한다", BLUE_LIGHT),
             ("다음 행동으로 순환", "승인 결과가 다시\n다음 추천을 부른다", BLUE)]
    for i, (t, d, ac) in enumerate(diffs):
        flat(sl, xs[i], top, w3, 3.1, fill=TINT, accent=ac)
        write(sl, xs[i] + 0.38, top + 0.5, w3 - 0.7, 0.45,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.5,
                          spacing_before=(0 if j == 0 else 2))) for j, s in enumerate(d.split("\n"))]
        write(sl, xs[i] + 0.38, top + 1.22, w3 - 0.7, 1.3, lines)
    keyline(sl, top + 3.42, "경쟁 제품에도 AI는 있다. 다른 것은 승인 결과가 다음 행동으로 순환하는 구조다")

    # ========================= 13 해결책 =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="기록이 초안이 되고, 승인이 다음 일정이 된다",
                    source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    w4, xs = grid(4, 0.3)
    sols = [("01", "미팅 기록 입력", "원문 · 음성 STT\n이미지 OCR"),
            ("02", "AI 초안 + 딜 승산", "딜별 귀속 · 13개 특성\n성사 확률"),
            ("03", "사람이 확정", "근거 대조 · 확정\n팀장 검토"),
            ("04", "다음 미팅 · 브리핑", "날짜 추천\n캘린더 승인")]
    for i, (n, t, d) in enumerate(sols):
        numcard(sl, xs[i], top, w4, 3.15, n, t, d, active=(i == 2), num_size=42, title_lines=1)
        if i < 3:
            arrow(sl, xs[i] + w4 + 0.04, top + 1.58, xs[i] + w4 + 0.26, top + 1.58,
                  color=BLUE_LIGHT, width=2)
    keyline(sl, top + 3.45, "04의 승인 결과가 다시 01로 — 모든 기록은 딜(sales_deal) 하나로 연결된다")

    # ========================= 14 기대효과 ① =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="파일럿에서 네 개 지표로 검증한다",
                    source="출처: 프로젝트 기획서 9장 · 파일럿 KPI 정의")
    w4, xs = grid(4, 0.3)
    kpis = [("보고서 작성시간", "활성 작업시간 중앙값"),
            ("당일 제출률", "당일 제출 ÷ 보고 대상"),
            ("필수정보 초회 완성률", "첫 제출 필수항목 충족"),
            ("팀장 검토시간", "검토 작업시간 중앙값")]
    for i, (t, d) in enumerate(kpis):
        placeholder(sl, xs[i], top, w4, 1.75, "[수치 입력 예정]", size=T_BODY)
        write(sl, xs[i], top + 1.98, w4, 0.4, [(t, dict(size=19, weight=600, color=INK, line=1.3))])
        write(sl, xs[i], top + 2.45, w4, 0.6,
              [(d, dict(size=T_BODY, weight=400, color=SUB, line=1.4))])
    keyline(sl, top + 3.6, "정량 수치는 파일럿에서 검증할 가설이며 실측값이 아니다")

    # ========================= 15 기대효과 ② =========================
    sl, top = slide(prs, p(), chapter=CH1, headline="보고서는 아래 단계의 확정본 위에 쌓인다",
                    source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    lw = 6.6
    levels = [("미팅 보고서", BLUE_PALE, INK), ("일일 업무보고", BLUE_LIGHT, WHITE),
              ("주간 업무보고", BLUE_500, WHITE), ("월간 업무보고", BLUE, WHITE)]
    for i, (t, fc, tc) in enumerate(levels):
        y = top + i * 1.0
        ins = i * 0.38
        sh = card(sl, M + ins, y, lw - ins * 2, 0.76, fill=fc, shadow=False, radius=0.1)
        tf = sh.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, t, size=T_BODY, weight=600, color=tc, line=1.2, align=C, first=True)
        if i < 3:
            arrow(sl, M + lw / 2, y + 0.79, M + lw / 2, y + 0.97, color=BLUE_LIGHT, width=2)
    note(sl, M, top + 4.05, lw, "각 단계는 아래 단계의 확정 제출본을 원천으로 사용한다", size=16)
    rx = M + lw + 0.6
    rw = CW - lw - 0.6
    label(sl, rx, top, "정성 효과", w=rw)
    items = [("자동 집계 · 검증", "흩어진 기록을 모아 숫자를 스스로 맞춘다"),
             ("승인 기반 반영", "확정된 내용만 CRM · 캘린더에 반영한다"),
             ("암묵지의 자산화", "개인의 노하우가 조직의 기록으로 남는다")]
    for i, (t, d) in enumerate(items):
        y = top + 0.52 + i * 1.32
        flat(sl, rx, y, rw, 1.12, fill=TINT)
        write(sl, rx + 0.3, y + 0.2, rw - 0.6, 0.35,
              [(t, dict(size=19, weight=600, color=BLUE_DEEP, line=1.25))])
        write(sl, rx + 0.3, y + 0.62, rw - 0.6, 0.4,
              [(d, dict(size=T_BODY, weight=400, color=SUB, line=1.35))])

    # ========================= 16 파트 디바이더 2 =========================
    divider(prs, "02", "설계와 구현", "그래서 어떻게 해결했는가", p())

    # ========================= 17 시나리오 ① =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="AI는 미팅을 준비하고 기록을 받아 적는다",
                    source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    label(sl, M, top, "사용자 시나리오 01 – 03")
    w3, xs = grid(3, 0.36)
    sc1 = [("01", "미팅 준비", "영업·계약관리 에이전트가\n다음 미팅을 추천하고\n사람이 캘린더에 승인한다"),
           ("02", "사전 브리핑", "과거 보고서와 자료를\nRAG로 찾아 근거와 함께\n브리핑을 만든다"),
           ("03", "미팅 후 입력", "미팅 원문을 직접 쓰거나\n음성 파일 STT · 이미지 OCR로\n기록을 올린다")]
    for i, (n, t, d) in enumerate(sc1):
        numcard(sl, xs[i], top + 0.48, w3, 3.4, n, t, d, num_size=42, title_lines=1)
    keyline(sl, top + 4.02, "여기까지는 모두 AI가 준비한 초안 — 아직 확정된 것은 없다")

    # ========================= 18 시나리오 ② =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="확정은 사람이 하고, 그 확정이 다음을 부른다",
                    source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    label(sl, M, top, "사용자 시나리오 04 – 06")
    w3, xs = grid(3, 0.36)
    sc2 = [("04", "보고서 확정", "사람이 확정해야\n저장된다\nAI 단독 저장은 없다"),
           ("05", "보고서 누적", "확정본이 RAG에 적재되고\n다음 미팅 추천\n트리거가 걸린다"),
           ("06", "기간 보고", "일일 · 주간 · 월간으로\n묶여 팀장 검토로\n올라간다")]
    for i, (n, t, d) in enumerate(sc2):
        numcard(sl, xs[i], top + 0.48, w3, 3.4, n, t, d, active=(i == 0), num_size=42, title_lines=1)
    keyline(sl, top + 4.02, "확정 결과가 다시 01 미팅 준비로 순환한다")

    # ========================= 19 시연 =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="한 번의 영업 사이클을 실제 화면으로 따라간다")
    vw = 7.4
    placeholder(sl, M, top, vw, vw * 9 / 16, "시연 영상 (16:9)", size=20)
    note(sl, M, top + vw * 9 / 16 + 0.14, vw, "[시연 동선 한 줄 입력 예정]", size=16)
    rx = M + vw + 0.5
    rw = CW - vw - 0.5
    label(sl, rx, top, "시연 체크포인트", w=rw)
    cps = ["자료실 업로드 → 요약", "캘린더 추천 승인", "AI 브리핑 확인",
           "미팅 원문 → 초안 생성", "보고서 확정 → 다음 추천"]
    for i, t in enumerate(cps):
        y = top + 0.52 + i * 0.76
        flat(sl, rx, y, rw, 0.64, fill=TINT)
        write(sl, rx + 0.26, y + 0.17, 0.4, 0.3,
              [(str(i + 1), dict(size=T_BODY, weight=700, color=BLUE, line=1.3))])
        write(sl, rx + 0.64, y + 0.17, rw - 0.88, 0.3,
              [(t, dict(size=T_BODY, weight=500, color=INK, line=1.3))])

    # ========================= 20 기술 스택 (한 장) =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="이 프로젝트를 만든 기술 조합",
                    source="출처: frontend/package.json · backend/uv.lock · 운영 설정값(2026-09-11 확인)")
    groups = [
        ("FRONTEND", [("react", "React 19"), ("typescript", "TypeScript"), ("vite", "Vite")]),
        ("BACKEND", [("fastapi", "FastAPI"), ("python", "Python 3.13"),
                     ("sqlalchemy", "SQLAlchemy")]),
        ("DATABASE", [("supabase", "Supabase"), ("postgresql", "PostgreSQL"),
                      (None, "pgvector")]),
        ("AI · ML", [("openai", "OpenAI API"), ("langchain", "LangChain"),
                     ("scikitlearn", "scikit-learn"), (None, "CatBoost"), (None, "DeepAgents")]),
        ("OCR", [(None, "PaddleOCR"), (None, "RunPod GPU")]),
        ("INFRA", [("docker", "Docker"), ("githubactions", "GitHub Actions"),
                   ("amazonwebservices", "AWS")]),
    ]
    gw, gxs = grid(6, 0.2)
    gh = 4.12
    for gi, (glabel, items) in enumerate(groups):
        x = gxs[gi]
        acc = (gi in (3, 4))
        flat(sl, x, top, gw, gh, fill=(TINT if acc else GRAY))
        write(sl, x, top + 0.24, gw, 0.3,
              [(glabel, dict(size=T_NOTE, weight=700, color=(BLUE if acc else BLUE_600),
                             line=1.2))], align=C)
        line(sl, x + 0.22, top + 0.64, x + gw - 0.22, top + 0.64, color=BORDER)
        for i, (slug, name) in enumerate(items):
            iy = top + 0.82 + i * 0.64
            if slug:
                put_icon(sl, slug, x + (gw - 0.34) / 2, iy, 0.34)
            else:
                monogram(sl, x + (gw - 0.34) / 2, iy, 0.34, name[0].upper())
            write(sl, x + 0.1, iy + 0.4, gw - 0.2, 0.3,
                  [(name, dict(size=15, weight=500, color=INK, line=1.2))], align=C)
    my = top + gh + 0.08
    flat(sl, M, my, CW, 0.38, fill=TINT2, accent=BLUE)
    write(sl, M + 0.34, my + 0.07, CW - 0.6, 0.28,
          [("LLM gpt-5.6-luna · STT gpt-4o-transcribe · OCR PaddleOCR (RunPod GPU) · "
            "provider adapter로 교체 가능",
            dict(size=15, weight=600, color=BLUE_DEEP, line=1.2))])

    # ========================= 21 시스템 아키텍처 (한 장) =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="입구는 하나, 오래 걸리는 AI 작업은 따로 뗐다",
                    source="출처: deploy/backend/deploy.sh · backend/app/services/agent_worker.py")
    lane_x, lane_w = 1.9, 5.5
    side_x, side_w = 7.75, 4.96
    rows = [
        (2.04, 0.50, "CLIENT", MUTED, GRAY, "사용자", None, None, None),
        (2.64, 0.72, "EDGE", BLUE_600, TINT, "CloudFront", "단일 HTTPS 진입점",
         "amazonwebservices", ("amazons3", "S3 (Private)", "React SPA 정적 호스팅")),
        (3.46, 0.72, "APP", BLUE_600, TINT, "EC2 · Nginx · FastAPI", "API 서버 · SSE 진행 상황",
         "amazonec2", (None, "OCR · STT 동기 호출",
                       "RunPod PaddleOCR · gpt-4o-transcribe")),
        (4.28, 1.14, "DATA", BLUE, TINT2, "Supabase",
         "PostgreSQL · pgvector · Auth · Storage\nagent_run 작업 큐", "supabase", None),
        (5.52, 0.72, "ASYNC", BLUE_500, TINT, "Agent Worker", "큐 폴링 · 별도 Docker 컨테이너",
         "docker", ("openai", "OpenAI gpt-5.6-luna", "보고서 · 미팅분석 생성")),
    ]
    prev = None
    for (y, h, zone, zc, fc, title, sub, ic, side) in rows:
        w_ = lane_w if zone != "CLIENT" else 2.6
        x_ = lane_x if zone != "CLIENT" else lane_x + (lane_w - 2.6) / 2
        flat(sl, x_, y, w_, h, fill=fc)
        write(sl, M, y + (h - 0.24) / 2, 1.2, 0.3,
              [(zone, dict(size=T_NOTE, weight=700, color=zc, line=1.2))])
        tx = x_ + (0.86 if ic else 0.32)
        if ic:
            put_icon(sl, ic, x_ + 0.32, y + (h - 0.34) / 2, 0.34)
        lines = [(title, dict(size=19, weight=700, color=INK, line=1.25))]
        if sub:
            for j, s in enumerate(sub.split("\n")):
                lines.append((s, dict(size=15, weight=400, color=SUB, line=1.3,
                                      spacing_before=(4 if j == 0 else 1))))
        write(sl, tx, y + (0.12 if sub else (h - 0.3) / 2), w_ - (tx - x_) - 0.25, h, lines)
        if prev is not None:
            arrow(sl, lane_x + lane_w / 2, prev, lane_x + lane_w / 2, y - 0.02,
                  color=BLUE_LIGHT, width=2)
        prev = y + h
        if side:
            s_ic, s_t, s_d = side
            flat(sl, side_x, y, side_w, h, fill=GRAY)
            if s_ic:
                put_icon(sl, s_ic, side_x + 0.3, y + (h - 0.32) / 2, 0.32)
            write(sl, side_x + 0.76, y + 0.12, side_w - 1.0, h,
                  [(s_t, dict(size=17, weight=600, color=INK, line=1.25)),
                   (s_d, dict(size=T_NOTE, weight=400, color=MUTED, line=1.3, spacing_before=4))])
            arrow(sl, lane_x + lane_w + 0.04, y + h / 2, side_x - 0.04, y + h / 2,
                  color=BLUE_PALE, width=1.5, dashed=True)
    dy = 6.32
    flat(sl, M, dy, CW, 0.42, fill=TINT2, accent=BLUE)
    write(sl, M + 0.34, dy + 0.08, CW - 0.6, 0.3,
          [("배포    GitHub Actions → OIDC → SSM 원격 실행    ·    무중단 blue / green",
            dict(size=16, weight=600, color=BLUE_DEEP, line=1.25))])

    # ========================= 22 에이전트 5종 개요 =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="다섯 에이전트가 판단 · 초안 · 제안을 나눠 맡는다",
                    source="출처: backend/app/agents/ (develop, 2026-09-17)")
    agents = [("미팅분석", "판단 생성", "미팅 원문에서 딜별 귀속과 성사 확률을 만든다", BLUE),
              ("영업 · 계약관리", "상태 허브", "위험 신호와 이력을 보고 다음 미팅을 제안한다", BLUE_LIGHT),
              ("보고서 작성", "실행 지원", "원문과 하위 확정 보고서로 초안을 쓴다", BLUE_PALE),
              ("일정관리", "실행 지원", "저장된 추천이 아직 유효한지 다시 판단한다", BLUE_PALE),
              ("자료요약", "실행 지원", "문서를 구조화 요약하고 RAG 청크로 적재한다", BLUE_PALE)]
    for i, (name, tag, role, ac) in enumerate(agents):
        y = top + i * 0.86
        flat(sl, M, y, CW, 0.72, fill=TINT, accent=ac)
        write(sl, M + 0.36, y + 0.19, 2.9, 0.36,
              [(name, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        pill(sl, M + 3.4, y + 0.2, 1.55, 0.34, tag, fill=WHITE, color=BLUE_600, size=T_NOTE)
        write(sl, M + 5.25, y + 0.21, CW - 5.6, 0.36,
              [(role, dict(size=T_BODY, weight=400, color=SUB, line=1.3))])

    # ========================= 23 에이전트 상세 ① =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="판단과 제안을 만드는 두 에이전트",
                    source="출처: backend/app/agents/ (develop, 2026-09-17)")
    w2, xs2 = grid(2, 0.5)

    def agent_card(x, y, w, h, name, tag, src, out, human, ac=BLUE):
        flat(sl, x, y, w, h, fill=TINT, accent=ac)
        write(sl, x + 0.38, y + 0.36, w - 0.7, 0.4,
              [(name, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        pill(sl, x + 0.38, y + 0.92, 1.55, 0.34, tag, fill=WHITE, color=BLUE_600, size=T_NOTE)
        for j, (k, v) in enumerate([("입력", src), ("출력", out), ("사람 확인", human)]):
            yy = y + 1.52 + j * 0.88
            write(sl, x + 0.38, yy, 1.3, 0.3,
                  [(k, dict(size=T_NOTE, weight=600, color=MUTED, line=1.25))])
            write(sl, x + 0.38, yy + 0.3, w - 0.7, 0.55,
                  [(v, dict(size=T_BODY, weight=(600 if j == 2 else 400),
                            color=(BLUE_DEEP if j == 2 else SUB), line=1.35))])

    agent_card(xs2[0], top, w2, 4.35, "미팅분석", "판단 생성", "미팅 원문",
               "딜별 귀속 · 13개 특성 · 성사 확률", "확정할 때 저장된다", BLUE)
    agent_card(xs2[1], top, w2, 4.35, "영업 · 계약관리", "상태 허브", "위험 신호 7종 · 이력",
               "다음 미팅 제안 · 근거 브리핑", "캘린더 승인", BLUE_LIGHT)

    # ========================= 24 에이전트 상세 ② =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="실행을 돕는 세 에이전트",
                    source="출처: backend/app/agents/ (develop, 2026-09-17)")
    w3, xs = grid(3, 0.36)
    agent_card(xs[0], top, w3, 4.35, "보고서 작성", "실행 지원", "원문 · 하위 확정 보고서",
               "작성 → 검토 → 수정 초안", "확정 · 팀장 검토", BLUE)
    agent_card(xs[1], top, w3, 4.35, "일정관리", "실행 지원", "저장된 추천 · 딜 상태",
               "유효 / 재생성 / 무시 판단", "승인 · 거절", BLUE_LIGHT)
    agent_card(xs[2], top, w3, 4.35, "자료요약", "실행 지원", "PDF · DOCX · 이미지",
               "구조화 요약 · RAG 청크", "원문 대조", BLUE_PALE)

    # ========================= 25 AI 기능 세 갈래 =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="ML · OCR · Agent는 서로 다른 일을 한다",
                    source="출처: backend/app/agents · OCR 고도화 전후 최종 결과보고서(2026-09-17)")
    tracks = [("ML", "딜 승산 예측", BLUE,
               ["미팅 원문", "딜 13개 특성", "Stacking_LR 모델", "성사 확률"]),
              ("OCR", "문서 텍스트 추출", BLUE_500,
               ["문서 · 이미지", "RunPod PDF OCR", "한글 누락 시\nPNG 재시도", "필드 추출\n고객 · 계약 등록"]),
              ("Agent", "초안과 제안", BLUE_LIGHT,
               ["업무 트리거", "agent_run 큐", "에이전트 초안", "사람이 확정"])]
    lw = 2.5
    fx = M + lw + 0.2
    fw = CW - lw - 0.2
    cw_ = (fw - 3 * 0.34) / 4
    rh = 1.38
    for i, (nm, role, ac, steps) in enumerate(tracks):
        y = top + i * (rh + 0.22)
        flat(sl, M, y, CW, rh, fill=TINT, accent=ac)
        write(sl, M + 0.36, y + 0.34, lw - 0.5, 0.36,
              [(nm, dict(size=22, weight=700, color=BLUE_DEEP, line=1.2))])
        write(sl, M + 0.36, y + 0.76, lw - 0.5, 0.3,
              [(role, dict(size=15, weight=400, color=SUB, line=1.25))])
        for j, st in enumerate(steps):
            cx = fx + j * (cw_ + 0.34)
            sh = card(sl, cx, y + 0.24, cw_, rh - 0.48, fill=WHITE, shadow=False, radius=0.1)
            tf = sh.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.1)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for k, s in enumerate(st.split("\n")):
                para(tf, s, size=16, weight=600, color=INK, line=1.3, align=C, first=(k == 0))
            if j < 3:
                arrow(sl, cx + cw_ + 0.05, y + rh / 2, cx + cw_ + 0.29, y + rh / 2,
                      color=BLUE_LIGHT, width=2)

    # ========================= 26 Agent Flow =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="트리거와 작업 큐로 잇고, 사람이 닫는다",
                    source="출처: backend/app/services/agent_runs.py · agent_worker.py")
    w4, xs = grid(4, 0.28)
    cols = [("업무 트리거", ["자료 업로드", "미팅 보고서 생성", "보고서 확정", "일정 등록",
                        "딜 단계 이동 외 3건"], GRAY, SUB),
            ("작업 큐", ["agent_run 큐", "PostgreSQL 기반", "부모 · 자식 실행", "재시도 최대 2회",
                      "lease 90초"], TINT, SUB),
            ("에이전트", ["미팅 내용분석", "보고서 작성", "딜 특성 + ML 예측", "계약관리 → 일정관리",
                      "브리핑 · 자료요약"], TINT2, INK),
            ("사람", ["보고서 확정", "추천 승인 · 일정 등록", "팀장 검토 · 반려"], TINT, SUB)]
    ch_h = 4.0
    for i, (lb, items, fc, tc) in enumerate(cols):
        x = xs[i]
        flat(sl, x, top, w4, ch_h, fill=fc, accent=(BLUE if i == 2 else None))
        write(sl, x + 0.26, top + 0.26, w4 - 0.5, 0.36,
              [(lb, dict(size=19, weight=700, color=(BLUE_DEEP if i == 2 else INK), line=1.25))])
        line(sl, x + 0.26, top + 0.76, x + w4 - 0.26, top + 0.76, color=BORDER)
        yy = top + 0.94
        for t in items:
            _bar(sl, x + 0.26, yy + 0.11, 0.08, 0.08, BLUE_LIGHT)
            write(sl, x + 0.46, yy, w4 - 0.66, 0.72,
                  [(t, dict(size=T_BODY, weight=(600 if i == 2 else 400), color=tc, line=1.3))])
            yy += 0.62
        if i < 3:
            arrow(sl, x + w4 + 0.04, top + ch_h / 2, x + w4 + 0.24, top + ch_h / 2,
                  color=BLUE_LIGHT, width=2)
    keyline(sl, top + ch_h + 0.16, "사람의 확정이 다시 트리거가 되어 처음으로 돌아간다")

    # ========================= 27 Supervisor 없음 =========================
    sl, top = slide(prs, p(), chapter=CH2, headline="전역 LLM Supervisor는 없다",
                    source="출처: backend/app/services/agent_runs.py · agent_worker.py")
    w3, xs = grid(3, 0.36)
    mech = [("무엇으로 잇는가", "agent_run 큐\n코드 dispatch\nAPI 트리거"),
            ("큐 규칙", "부모 · 자식 실행 구조\n재시도 최대 2회\nlease 90초"),
            ("DeepAgents 범위", "보고서 작성 내부에서만\n사용한다\n전체 조율에는 쓰지 않는다")]
    for i, (t, d) in enumerate(mech):
        flat(sl, xs[i], top, w3, 3.5, fill=(TINT2 if i == 2 else TINT),
             accent=(BLUE if i == 2 else None))
        write(sl, xs[i] + 0.38, top + 0.42, w3 - 0.7, 0.4,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.6,
                          spacing_before=(0 if j == 0 else 3))) for j, s in enumerate(d.split("\n"))]
        write(sl, xs[i] + 0.38, top + 1.15, w3 - 0.7, 2.2, lines)
    keyline(sl, top + 3.8, "연결은 큐와 코드로 — 판단은 에이전트가, 확정은 사람이 한다")

    # ========================= 28 파트 디바이더 3 =========================
    divider(prs, "03", "검증과 확장", "제대로 동작하는지 어떻게 확인했는가", p())

    # ---------- 공통 헬퍼 ----------
    def kpi_row(sl, y, items, *, featured=None, h=1.78):
        n = len(items)
        w_, xs_ = grid(n, 0.3)
        for i, (v, lb, sb) in enumerate(items):
            flat(sl, xs_[i], y, w_, h, fill=(TINT2 if featured == i else TINT),
                 accent=(BLUE if featured == i else None))
            kpi(sl, xs_[i] + 0.34, y + 0.26, w_ - 0.68, v, lb, sb,
                color=(BLUE if featured == i else INK))
        return w_, xs_

    def limit_cards(sl, y, items, *, h=3.6):
        w_, xs_ = grid(len(items), 0.36)
        for i, (t, d) in enumerate(items):
            flat(sl, xs_[i], y, w_, h, fill=GRAY)
            write(sl, xs_[i] + 0.34, y + 0.4, w_ - 0.66, 0.4,
                  [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
            lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.55,
                              spacing_before=(0 if j == 0 else 2)))
                     for j, s in enumerate(d.split("\n"))]
            write(sl, xs_[i] + 0.34, y + 1.1, w_ - 0.66, h - 1.2, lines)

    def two_notes(sl, y, h, blocks, *, gap=0.44):
        w_, xs_ = grid(len(blocks), 0.4)
        for i, (t, items) in enumerate(blocks):
            flat(sl, xs_[i], y, w_, h, fill=GRAY)
            write(sl, xs_[i] + 0.34, y + 0.24, w_ - 0.66, 0.34,
                  [(t, dict(size=19, weight=700, color=INK, line=1.25))])
            bullets(sl, xs_[i] + 0.34, y + 0.72, w_ - 0.68, items, gap=gap, size=16, marker=MUTED)

    # ========================= 29 평가 개요 =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="에이전트마다 다른 방법으로 검증했다",
                    source="출처: docs/eval · output/evals · OCR 고도화 전후 최종 결과보고서(2026-09-17)")
    rows_ = [("미팅분석", "합성 골든셋 정답표 대조 + 딜 승산 ML 30회 반복"),
             ("보고서 작성", "합성 53건 LLM Judge 절대 평가 + 개선 전후 쌍대 비교"),
             ("자료요약 · RAG", "합성 문서 7개 · QRA 자동 골든셋 10건 LLM Judge"),
             ("OCR", "문서 200건 선정 필드 정답값 대조 · 1차와 최종 재시험 비교"),
             ("영업 · 계약관리", "[평가 방법 입력 예정]")]
    for i, (t, d) in enumerate(rows_):
        y = top + i * 0.92
        done = i < 4
        flat(sl, M, y, CW, 0.78, fill=(TINT if done else GRAY),
             accent=(BLUE_LIGHT if done else None))
        write(sl, M + 0.36, y + 0.22, 4.4, 0.36,
              [(t, dict(size=T_CARD, weight=700, color=(INK if done else MUTED), line=1.25))])
        write(sl, M + 5.0, y + 0.24, CW - 5.35, 0.36,
              [(d, dict(size=T_BODY, weight=400, color=(SUB if done else MUTED), line=1.3))])

    # ========================= 30 미팅분석 ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="딜 귀속 84.7%, 특성 추출 74.8%가 다음 과제",
                    source="출처: 미팅 내용분석 구조 보고서 (2026-08-31 골든셋 · 과거 코드 기준)")
    kpi_row(sl, top, [("84.7%", "근거 · 딜 일치", "94 / 111 구간"),
                      ("74.8%", "딜 특성 필드 정확도", "350 / 468 필드"),
                      ("82.6%", "필수 사실 보존", "38 / 46"),
                      ("0 / 36", "13개 특성 전부 정답", "가장 큰 과제")], featured=3)
    label(sl, M, top + 2.06, "평가 방법")
    bullets(sl, M, top + 2.54, CW, ["합성 미팅 12회 · 딜 36개로 만든 골든셋 (2026-08-31)",
                                    "AI가 의미를 검수한 결과이며 사람 판정이 아니다"], gap=0.54)
    keyline(sl, top + 3.72, "13개 특성을 한 건도 빠짐없이 맞춘 사례는 36건 중 0건 — 가장 먼저 개선할 지점")

    # ========================= 31 미팅분석 ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="딜 승산 모델은 Stacking으로 AUC 0.778까지 올랐다",
                    source="출처: 머신러닝 · 딥러닝 학습결과서 (30회 평균)")
    cwid = 8.0
    label(sl, M, top, "단계별 성능 — 공개 영문 B2B 448건 · 30회 반복 평균", w=cwid)
    add_chart(sl, COL, M - 0.2, top + 0.46, cwid, 3.55,
              ["RF 기준선", "RF 튜닝", "Stacking_LR (최종)"],
              [("Accuracy", [0.6958, 0.7338, 0.7360]), ("ROC-AUC", [0.7409, 0.7756, 0.7783])],
              label_fmt="0.000", gap=80, y_min=0.6, y_max=0.84, colors=(BLUE_PALE, BLUE),
              axis_size=T_NOTE, label_size=T_NOTE)
    rx = M + cwid + 0.5
    rw = CW - cwid - 0.5
    kpi(sl, rx, top, rw, "+0.037", "ROC-AUC 상승폭", "0.7409 → 0.7783", tint=True, h=1.85)
    kpi(sl, rx, top + 2.02, rw, "0.1880", "Brier 점수", "0.2097 → 0.1880 (낮을수록 좋음)",
        color=BLUE_500, tint=True, h=1.95)

    # ========================= 32 미팅분석 ③ =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="미팅분석 수치를 읽을 때 유의할 점",
                    source="출처: 미팅 내용분석 구조 보고서 · 머신러닝 · 딥러닝 학습결과서")
    limit_cards(sl, top, [("합성 데이터", "합성 미팅 12회 · 딜 36개\n2026-08-31 골든셋\n과거 코드 기준"),
                          ("사람 판정 아님", "AI가 의미를 검수했다\n사람 평가자가\n채점한 결과가 아니다"),
                          ("영문 공개 데이터", "딜 승산 ML은\n공개 영문 B2B 448건 기준\n한국어 실제 성능은 별도")],
                h=3.6)
    keyline(sl, top + 3.86, "실제 한국어 영업 데이터에서의 성능은 아직 측정하지 않았다")

    # ========================= 33 보고서 ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="구조를 바꾸자 53건 모두 생성, 토큰은 54% 감소",
                    source="출처: docs/eval/01_보고서작성에이전트 (Judge gpt-5.6-luna, 2026-09-08)")
    kpi_row(sl, top, [("47 → 53", "유효 답안", "53건 중"),
                      ("18 → 0", "외부 재시도", "생성 실패 6 → 0"),
                      ("−54.4%", "생성 토큰", "1,001만 → 457만"),
                      ("29 / 47", "쌍대 우세", "개선 전 우세 1건")], featured=2)
    label(sl, M, top + 2.06, "평가 방법")
    bullets(sl, M, top + 2.54, CW, ["루브릭 5개로 절대 채점하고, 순서를 익명 교차해 쌍대 비교했다",
                                    "인용은 기계로 검증하고, 기간 보고서는 하위 확정본을 참조했다"], gap=0.54)
    keyline(sl, top + 3.72, "실패하던 6건이 사라지고, 같은 결과를 절반 이하의 토큰으로 만든다")

    # ========================= 34 보고서 ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="주간 · 월간 보고서 점수가 가장 크게 올랐다",
                    source="출처: docs/eval/01_보고서작성에이전트 (Judge gpt-5.6-luna, 2026-09-08)")
    cwid = 8.0
    label(sl, M, top, "보고서 유형별 LLM Judge 점수 — 100점 만점 · 공통 43건", w=cwid)
    add_chart(sl, COL, M - 0.2, top + 0.46, cwid, 3.55, ["미팅", "일일", "주간", "월간"],
              [("개선 전", [97.74, 86.88, 72.50, 66.25]), ("개선 후", [96.57, 95.16, 87.08, 76.25])],
              label_fmt="0.0", gap=70, y_min=60, y_max=108, colors=(BLUE_PALE, BLUE),
              axis_size=T_NOTE, label_size=T_NOTE)
    rx = M + cwid + 0.5
    rw = CW - cwid - 0.5
    kpi(sl, rx, top, rw, "+14.6", "주간 보고서 상승", "72.50 → 87.08", tint=True, h=1.85)
    kpi(sl, rx, top + 2.02, rw, "−1.2", "미팅 보고서 하락", "97.74 → 96.57",
        color=MUTED, tint=True, h=1.95)

    # ========================= 35 보고서 ③ =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="보고서 수치를 읽을 때 유의할 점",
                    source="출처: docs/eval/01_보고서작성에이전트 (2026-09-08)")
    limit_cards(sl, top, [("표본이 작다", "시나리오 1개\n월간 보고서는 1건뿐"),
                          ("같은 모델이 채점", "생성 모델과 Judge 모델이\n동일하다 (gpt-5.6-luna)"),
                          ("떨어진 지표도 있다", "미팅 보고서 점수 하락\n사실 정확성 3.84 → 3.74")],
                h=3.6)
    keyline(sl, top + 3.86, "개선은 기간 보고서에 집중되었고, 미팅 보고서의 사실 정확성은 오히려 떨어졌다")

    # ========================= 36 영업 · 계약관리 (레이아웃) =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="[영업 · 계약관리 평가 결론 입력 예정]")
    w4, xs = grid(4, 0.3)
    for i in range(4):
        placeholder(sl, xs[i], top, w4, 1.7, "[지표 입력 예정]")
    placeholder(sl, M, top + 1.95, 8.0, 2.63, "[차트 또는 평가 도식 입력 예정]")
    placeholder(sl, M + 8.3, top + 1.95, CW - 8.3, 2.63, "[평가 데이터 · 방법 · 한계 입력 예정]")

    # ========================= 37 일정 에이전트 =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="일정 추천은 보고서에 적힌 날짜에서 출발한다",
                    source="출처: backend/app/agents/contract_management.py · schedule_management.py")
    label(sl, M, top, "현재 동작")
    w4, xs = grid(4, 0.28)
    steps = [("01", "보고서 확정", "확정된 미팅 보고서가\n트리거가 된다"),
             ("02", "합의 날짜 추출", "보고서에 적힌 다음 만남\n날짜를 그대로 쓴다"),
             ("03", "없으면 이력 판단", "미팅 간격 · 요일과 딜\n맥락으로 날짜를 고른다"),
             ("04", "사람이 캘린더 승인", "담당자가 승인하거나\n거절한다")]
    sy = top + 0.38
    for i, (n, t, d) in enumerate(steps):
        act = i == 3
        flat(sl, xs[i], sy, w4, 1.72, fill=(TINT2 if act else TINT), accent=(BLUE if act else None))
        write(sl, xs[i] + 0.3, sy + 0.14, 0.9, 0.44,
              [(n, dict(size=26, weight=700, color=(BLUE if act else BLUE_LIGHT), line=1.2))])
        write(sl, xs[i] + 0.3, sy + 0.64, w4 - 0.6, 0.36,
              [(t, dict(size=19, weight=700, color=INK, line=1.25))])
        lines = [(s, dict(size=16, weight=400, color=SUB, line=1.4,
                          spacing_before=(0 if j == 0 else 1))) for j, s in enumerate(d.split("\n"))]
        write(sl, xs[i] + 0.3, sy + 1.02, w4 - 0.6, 0.62, lines)
        if i < 3:
            arrow(sl, xs[i] + w4 + 0.03, sy + 0.86, xs[i] + w4 + 0.25, sy + 0.86,
                  color=BLUE_LIGHT, width=2)
    label(sl, M, top + 2.32, "승인한 뒤")
    two_notes(sl, top + 2.72, 1.5,
              [("일정관리 에이전트의 재판단",
                ["유효 / 재생성 / 무시 중 하나로 판단",
                 "날짜는 직접 만들지 않고 계약관리에 요청"]),
               ("지금 하지 않는 것",
                ["외부 캘린더 연동은 없다",
                 "반복 일정 패턴을 자동 규칙화하지 않는다"])],
              gap=0.38)
    note(sl, M, top + 4.32, CW,
         "규칙 기반 자율 스케줄링이 아니라, 보고서에 적힌 날짜를 우선 쓰는 단순한 방식이다", size=16)

    # ========================= 38 문서요약 ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="문서요약 자동 평가 종합 점수는 0.846이다",
                    source="출처: docs/document-summary-evaluation.md · output/evals/document-summary-rageval")
    cwid = 8.6
    label(sl, M, top, "지표별 점수 (0–1 · 높을수록 좋음)", w=cwid)
    add_chart(sl, BAR, M - 0.15, top + 0.46, cwid, 4.0,
              ["overall", "completeness", "groundedness", "hallucination",
               "retrieval_relevance", "irrelevance"],
              [("점수", [0.846, 0.828, 0.914, 0.942, 0.892, 0.901])],
              label_fmt="0.000", gap=55, y_min=0.0, y_max=1.0,
              axis_size=T_NOTE, label_size=T_NOTE)
    rx = M + cwid + 0.5
    rw = CW - cwid - 0.5
    kpi(sl, rx, top, rw, "0.846", "overall", "자동 골든셋 10건 종합", tint=True, h=1.85)
    kpi(sl, rx, top + 2.02, rw, "0.828", "completeness", "가장 낮은 지표",
        color=BLUE_500, tint=True, h=1.95)

    # ========================= 39 문서요약 ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="합성 문서에서 골든셋을 만들어 자동 채점했다",
                    source="출처: docs/document-summary-evaluation.md")
    lw = 6.6
    label(sl, M, top, "평가 흐름", w=lw)
    fsteps = [("01", "합성 문서 7개"), ("02", "QRA 자동 골든셋 10건"),
              ("03", "요약 6 · 검색 4 LLM Judge 채점")]
    for i, (n, t) in enumerate(fsteps):
        y = top + 0.52 + i * 1.14
        flat(sl, M, y, lw, 0.9, fill=TINT, accent=(BLUE if i == 2 else None))
        write(sl, M + 0.34, y + 0.25, 0.7, 0.4,
              [(n, dict(size=T_CARD, weight=700, color=BLUE, line=1.2))])
        write(sl, M + 1.1, y + 0.27, lw - 1.4, 0.4,
              [(t, dict(size=T_BODY, weight=500, color=INK, line=1.3))])
        if i < 2:
            arrow(sl, M + lw / 2, y + 0.93, M + lw / 2, y + 1.11, color=BLUE_LIGHT, width=2)
    rx = M + lw + 0.6
    rw = CW - lw - 0.6
    label(sl, rx, top, "읽을 때 주의", w=rw, color=MUTED)
    flat(sl, rx, top + 0.52, rw, 3.4, fill=GRAY)
    bullets(sl, rx + 0.34, top + 0.94, rw - 0.68,
            ["요약 6건은 0.91–0.98 구간", "사람 검수는 pending 상태", "모두 합성 문서 기준 결과"],
            gap=0.9, marker=MUTED)

    # ========================= 40 OCR ① 결과 =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="조건부 PNG 재시도로 핵심 문서 일치율 98.8%",
                    source="출처: OCR 고도화 전후 최종 결과보고서 (2026-09-17 RunPod 재시험)")
    kpi_row(sl, top, [("98.8%", "핵심 구조화 문서", "1,136 / 1,150 · +33.4%p"),
                      ("98.5%", "계약서", "197 / 200 · +73.5%p"),
                      ("98.5%", "발주서", "739 / 750 · +31.6%p"),
                      ("100%", "견적서", "200 / 200 · 1차부터 유지")], featured=0)
    label(sl, M, top + 2.08, "문서 유형별 선정 필드 일치율 (%)")
    ch = add_chart(sl, COL, M - 0.2, top + 2.5, CW + 0.4, 2.08,
                   ["견적서", "계약서", "발주서", "상품설명서"],
                   [("1차 테스트", [100.0, 25.0, 66.9, 84.0]),
                    ("최종 테스트", [100.0, 98.5, 98.5, 86.0])],
                   label_fmt="0.0", gap=70, y_max=120, colors=(BLUE_PALE, BLUE),
                   axis_size=T_NOTE, label_size=T_NOTE)
    ch.value_axis.major_unit = 30

    # ========================= 41 OCR ② 방법과 한계 =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="스캔 PDF만 골라 이미지 OCR로 다시 읽는다",
                    source="출처: OCR 고도화 전후 최종 결과보고서 (2026-09-17)")
    label(sl, M, top, "고도화 3단계")
    w3, xs = grid(3, 0.36)
    oc = [("1차", "원본 PDF를 RunPod\nPDF OCR로 처리한다"),
          ("조건부 재시도", "한글 필수값 누락 또는 작업\n실패일 때만 모든 페이지를\nPNG로 변환해 다시 읽는다"),
          ("결과 선택", "발주서 공급자 주소는 PDF\n결과를 남기고, 나머지 필드는\nPNG 결과를 쓴다")]
    sy = top + 0.42
    for i, (t, d) in enumerate(oc):
        act = i == 1
        flat(sl, xs[i], sy, w3, 1.95, fill=(TINT2 if act else TINT), accent=(BLUE if act else None))
        write(sl, xs[i] + 0.34, sy + 0.26, w3 - 0.66, 0.36,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        lines = [(s, dict(size=16, weight=400, color=SUB, line=1.45,
                          spacing_before=(0 if j == 0 else 1))) for j, s in enumerate(d.split("\n"))]
        write(sl, xs[i] + 0.34, sy + 0.78, w3 - 0.66, 1.0, lines)
        if i < 2:
            arrow(sl, xs[i] + w3 + 0.05, sy + 0.92, xs[i] + w3 + 0.31, sy + 0.92,
                  color=BLUE_LIGHT, width=2)
    two_notes(sl, top + 2.48, 2.1,
              [("별도 관리 대상",
                ["명함 OCR 필드 일치 127 / 134 = 94.8%",
                 "사업자등록증은 정답지가 없어 비교 제외",
                 "엑셀 등록은 OCR이 아닌 파일 파싱"]),
               ("읽을 때 주의",
                ["전사 정확도 · CER / WER 미포함",
                 "손글씨 오류율 45%는 별도 과제",
                 "사양 값 자동 등록은 사람 검수 필요"])])

    # ========================= 42 RAG ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="요약은 고르게 높지만 검색 답변은 편차가 크다",
                    source="출처: output/evals/document-summary-rageval/llm_judge_results.json")
    label(sl, M, top, "케이스별 overall 점수 (0–1) · 요약 6건 · 검색 4건")
    add_chart(sl, COL, M - 0.2, top + 0.46, CW + 0.4, 3.4,
              ["요약1", "요약2", "요약3", "요약4", "요약5", "요약6", "검색1", "검색2", "검색3", "검색4"],
              [("overall", [0.97, 0.91, 0.96, 0.98, 0.93, 0.96, 0.25, 0.78, 0.90, 0.82])],
              label_fmt="0.00", gap=45, y_max=1.05, axis_size=T_NOTE, label_size=T_NOTE)
    keyline(sl, top + 3.96, "검색1 한 건이 0.25 — 최신 자료 대신 옛 계약서를 참조했다")

    # ========================= 43 RAG ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="RAG 답변 4건 중 1건은 옛 계약서 날짜를 답했다",
                    source="출처: output/evals/document-summary-rageval/llm_judge_results.json")
    lw = 7.0
    flat(sl, M, top, lw, 3.8, fill=TINT2, accent=BLUE)
    write(sl, M + 0.4, top + 0.38, lw - 0.8, 0.4,
          [("검색1 — overall 0.25", dict(size=T_CARD, weight=700, color=BLUE_DEEP, line=1.25))])
    for i, (k, v) in enumerate([("질문", "현재 기준 설치 예정일"),
                                ("기대한 답", "2026-10-08 (추가 자료)"),
                                ("실제 답변", "2026-09-24 (기존 계약서)")]):
        y = top + 1.06 + i * 0.76
        write(sl, M + 0.4, y, 1.7, 0.36, [(k, dict(size=T_NOTE, weight=600, color=MUTED, line=1.25))])
        write(sl, M + 2.2, y - 0.04, lw - 2.6, 0.4,
              [(v, dict(size=T_BODY, weight=500, color=INK, line=1.3))])
    write(sl, M + 0.4, top + 3.25, lw - 0.8, 0.4,
          [("→ 최신 자료 대신 옛 계약서를 참조했다",
            dict(size=T_BODY, weight=600, color=BLUE, line=1.3))])
    rx = M + lw + 0.6
    rw = CW - lw - 0.6
    label(sl, rx, top, "읽을 때 주의", w=rw, color=MUTED)
    flat(sl, rx, top + 0.52, rw, 3.28, fill=GRAY)
    bullets(sl, rx + 0.34, top + 0.84, rw - 0.68,
            ["평가 스크립트의 검색은 단어 겹침 방식이다",
             "운영은 키워드 + pgvector를 RRF로 병합하는 하이브리드 검색이다",
             "보고서 RAG · 브리핑 근거는 별도 품질 평가가 없다"],
            gap=0.98, marker=MUTED)

    # ========================= 44 실패와 개선 ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="[실패와 개선 사례 1 결론 입력 예정]")
    w4, xs = grid(4, 0.3)
    for i, lb in enumerate(["처음 접근", "무엇이 문제였나", "어떻게 바꿨나", "결과"]):
        write(sl, xs[i], top, w4, 0.36, [(lb, dict(size=19, weight=700, color=BLUE_600, line=1.25))])
        placeholder(sl, xs[i], top + 0.5, w4, 4.08, PH)

    # ========================= 45 실패와 개선 ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="[실패와 개선 사례 2 결론 입력 예정]")
    for i, lb in enumerate(["처음 접근", "무엇이 문제였나", "어떻게 바꿨나", "결과"]):
        write(sl, xs[i], top, w4, 0.36, [(lb, dict(size=19, weight=700, color=BLUE_600, line=1.25))])
        placeholder(sl, xs[i], top + 0.5, w4, 4.08, PH)

    # ========================= 46 보완점 =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="[보완점 결론 입력 예정]")
    w3, xs = grid(3, 0.36)
    for i in range(3):
        write(sl, xs[i], top, w3, 0.36,
              [("한계 %d" % (i + 1), dict(size=19, weight=700, color=BLUE_600, line=1.25))])
        placeholder(sl, xs[i], top + 0.5, w3, 3.1, "[영역 · 내용 · 영향 입력 예정]")
    placeholder(sl, M, top + 3.78, CW, 0.8, "[한 줄 정리 입력 예정]")

    # ========================= 47 향후 계획 Roadmap =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="다음 단계는 네 가지다")
    label(sl, M, top, "향후 계획 — 현재 구현된 기능이 아니다")
    w2, xs2 = grid(2, 0.4)
    plans = [("01", "일정 에이전트 고도화",
              "축적된 일정 데이터로\n반복 업무 패턴을 규칙화해\nRule-based 추천으로 확장"),
             ("02", "ML 모델 성능 고도화",
              "데이터와 모델 개선을 통한\n예측 성능 향상"),
             ("03", "팀장 화면 개선",
              "더 쉽고 빠른 업무 현황 확인을\n위한 UI/UX 개선\n가독성 · 현황 파악 · 접근성"),
             ("04", "매출 분석 기능 확장",
              "계약금 중심 관리에서\n매출 현황 · 추이 분석으로 확장")]
    for i, (n, t, d) in enumerate(plans):
        x = xs2[i % 2]
        y = top + 0.42 + (i // 2) * 2.21
        flat(sl, x, y, w2, 1.95, fill=(TINT2 if i == 0 else TINT),
             accent=(BLUE if i == 0 else None))
        ghost(sl, x + 0.36, y + 0.26, 1.0, n, size=40,
              color=(BLUE if i == 0 else BLUE_LIGHT))
        write(sl, x + 1.5, y + 0.3, w2 - 1.86, 0.4,
              [(t, dict(size=20, weight=700, color=INK, line=1.3))])
        dl = [(s, dict(size=17, weight=400, color=SUB, line=1.45,
                       spacing_before=(0 if j == 0 else 1))) for j, s in enumerate(d.split("\n"))]
        write(sl, x + 1.5, y + 0.84, w2 - 1.86, 1.05, dl)

    # ========================= 48 결론 ① =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="담당자의 머릿속을 조직의 기록으로 남긴다",
                    source="출처: SalesLuv 데이터 구조(develop, 2026-09-17) · 인수인계 질문 11개 영역 (1 / 2)")
    w3, xs = grid(3, 0.36)
    areas1 = [("고객 관계 맥락", "회사 · 담당자 · 부서"),
              ("접점 히스토리", "일정 · 활동 · 보고서"),
              ("고객의 실제 발언", "원문 · STT · 첨부"),
              ("발언의 대상 · 범위", "공통 · 딜별 + 근거 ID"),
              ("구매 판단 신호", "권한 · 경쟁사 등 13개 특성"),
              ("딜 진행 맥락", "단계 · 금액 · 계약 · 납기")]
    for i, (t, d) in enumerate(areas1):
        x = xs[i % 3]
        y = top + (i // 3) * 2.15
        flat(sl, x, y, w3, 1.85, fill=TINT)
        write(sl, x + 0.34, y + 0.42, w3 - 0.66, 0.4,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        write(sl, x + 0.34, y + 1.02, w3 - 0.66, 0.6,
              [(d, dict(size=T_BODY, weight=400, color=SUB, line=1.4))])

    # ========================= 49 결론 ② =========================
    sl, top = slide(prs, p(), chapter=CH3, headline="기록이 남으면 사람이 바뀌어도 맥락은 남는다",
                    source="출처: SalesLuv 데이터 구조(develop, 2026-09-17) · 인수인계 질문 11개 영역 (2 / 2)")
    areas2 = [("과거 맥락", "이전 확정 보고서"),
              ("이슈 · 불만 맥락", "불만 · 긴급도 · 처리"),
              ("위험과 다음 행동", "만료 · 지연 · 추천"),
              ("문서 지식", "원문 추출 · 요약 · RAG"),
              ("판단의 책임 · 이력", "작성자 · revision · 승인")]
    for i, (t, d) in enumerate(areas2):
        x = xs[i % 3]
        y = top + (i // 3) * 2.15
        flat(sl, x, y, w3, 1.85, fill=TINT)
        write(sl, x + 0.34, y + 0.42, w3 - 0.66, 0.4,
              [(t, dict(size=T_CARD, weight=700, color=INK, line=1.25))])
        write(sl, x + 0.34, y + 1.02, w3 - 0.66, 0.6,
              [(d, dict(size=T_BODY, weight=400, color=SUB, line=1.4))])
    gy = top + 2.15
    flat(sl, xs[2], gy, w3, 1.85, fill=DARK)
    write(sl, xs[2] + 0.34, gy + 0.46, w3 - 0.66, 1.1,
          [("AI는 구조화 · 요약 · 제안까지", dict(size=T_BODY, weight=700, color=WHITE, line=1.45)),
           ("확정은 사람의 승인으로", dict(size=T_BODY, weight=700, color=BLUE_LIGHT, line=1.45,
                                  spacing_before=8))])

    # ========================= 50 Q&A =========================
    sl = blank(prs, dark=True)
    sl.shapes.add_picture(LOGO_WHITE, Inches(SW - M - LOGO_W), Inches(EYE_Y - 0.02),
                          Inches(LOGO_W), Inches(LOGO_H))
    write(sl, M, 2.9, CW, 1.3, [("Q & A", dict(size=64, weight=700, color=WHITE, line=1.15))], align=C)
    _bar(sl, SW / 2 - 1.0, 4.45, 2.0, 0.075, BLUE)
    write(sl, M, 4.8, CW, 0.6,
          [("감사합니다 · Thank You", dict(size=21, weight=500, color=DARK_SUB, line=1.4))], align=C)
    write(sl, SW - M - 1.0, 6.83, 1.0, 0.3,
          [("%02d" % p(), dict(size=16, weight=600, color=SUB, line=1.25))], align=R)

    # ========================= 51 부록 ERD =========================
    sl, top = slide(prs, p(), chapter=APX, headline="모든 데이터는 딜을 중심으로 연결된다",
                    source="출처: backend/app/models/ (develop, 2026-09-17)")
    gw = 3.75
    gxs2 = [M, M + 4.17, M + 8.34]
    gh = 1.6
    top_y, bot_y = top, 5.02
    hub_w, hub_h = 3.4, 1.0
    hub_x, hub_y = SW / 2 - hub_w / 2, 3.83
    groups = [("조직 · 권한", "team · member", gxs2[0], top_y),
              ("고객", "customer_company\ncustomer_contact", gxs2[1], top_y),
              ("영업 실행", "sales_pipeline(+stage)\nproduct · sales_deal_item\npurchase_order(+item)",
               gxs2[2], top_y),
              ("활동 · 보고", "activity · report\nreport_deal · report_submission", gxs2[0], bot_y),
              ("AI 실행", "agent_run\ncontract_next_meeting_suggestion", gxs2[1], bot_y),
              ("문서 · C/S", "document · file\ndocument_chunk · support_request", gxs2[2], bot_y)]
    bus = hub_y + hub_h / 2
    line(sl, gxs2[0] + gw / 2, top_y + gh, gxs2[0] + gw / 2, bot_y, color=BLUE_PALE, width=1.5)
    line(sl, gxs2[2] + gw / 2, top_y + gh, gxs2[2] + gw / 2, bot_y, color=BLUE_PALE, width=1.5)
    line(sl, gxs2[0] + gw / 2, bus, hub_x, bus, color=BLUE_PALE, width=1.5)
    line(sl, hub_x + hub_w, bus, gxs2[2] + gw / 2, bus, color=BLUE_PALE, width=1.5)
    line(sl, gxs2[1] + gw / 2, top_y + gh, gxs2[1] + gw / 2, hub_y, color=BLUE_PALE, width=1.5)
    line(sl, gxs2[1] + gw / 2, hub_y + hub_h, gxs2[1] + gw / 2, bot_y, color=BLUE_PALE, width=1.5)
    hub = card(sl, hub_x, hub_y, hub_w, hub_h, fill=BLUE, shadow=False, radius=0.12)
    tf = hub.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "sales_deal", size=T_CARD, weight=700, color=WHITE, line=1.2, align=C, first=True)
    para(tf, "영업 딜", size=T_NOTE, weight=400, color=BLUE_PALE, line=1.3, align=C)
    for lb, items, x, y in groups:
        flat(sl, x, y, gw, gh, fill=TINT)
        write(sl, x + 0.28, y + 0.22, gw - 0.55, 0.34,
              [(lb, dict(size=17, weight=700, color=BLUE_600, line=1.25))])
        lines = [(t, dict(size=T_NOTE, weight=400, color=SUB, line=1.45,
                          spacing_before=(0 if j == 0 else 2)))
                 for j, t in enumerate(items.split("\n"))]
        write(sl, x + 0.28, y + 0.7, gw - 0.55, 0.8, lines)

    # ========================= 52 부록 WBS =========================
    sl, top = slide(prs, p(), chapter=APX, headline="7주 동안 기획에서 배포 검증까지 진행했다",
                    source="출처: 프로젝트 WBS")
    weeks = [("1주", "요구사항 분석\n기획"), ("2주", "DB 기반 구축\n화면 설계"),
             ("3주", "FE · BE 개발\n중간 발표"), ("4주", "데이터 전처리\nAI 모델링"),
             ("5주", "모델 평가\n기능 통합"), ("6주", "통합 테스트\n배포 검증"),
             ("7주", "산출물 검수\n최종 발표")]
    n = len(weeks)
    w_, xs7 = grid(n, 0.2)
    ry = top + 1.96
    line(sl, M + w_ / 2, ry, M + CW - w_ / 2, ry, color=BORDER, width=2)
    for i, (wk, t) in enumerate(weeks):
        x = xs7[i]
        write(sl, x, ry - 0.8, w_, 0.4, [(wk, dict(size=T_CARD, weight=700, color=INK, line=1.2))],
              align=C)
        circle(sl, x + w_ / 2 - 0.11, ry - 0.11, 0.22, "",
               fill=(BLUE if i >= n - 2 else BLUE_PALE))
        lines = [(s, dict(size=T_BODY, weight=400, color=SUB, line=1.45,
                          spacing_before=(0 if j == 0 else 2))) for j, s in enumerate(t.split("\n"))]
        write(sl, x, ry + 0.38, w_, 1.2, lines, align=C)

    # ========================= 53 부록 ML 성능 상세 =========================
    sl, top = slide(prs, p(), chapter=APX, headline="딜 승산 모델 단계별 성능 상세",
                    source="출처: 머신러닝 · 딥러닝 학습결과서 (공개 영문 B2B 448건 · 30회 반복 평균)")
    rows = [["단계", "Accuracy", "ROC-AUC", "Brier"],
            ["RF 기준선", "0.6958", "0.7409", "0.2097"],
            ["RF 튜닝", "0.7338", "0.7756", "–"],
            ["Stacking_LR (최종)", "0.7360", "0.7783", "0.1880"]]
    table(sl, M, top, CW, 3.2, rows, col_w=[4.2, 2.631, 2.631, 2.631], highlight_row=3)
    note(sl, M, top + 3.4, CW,
         "Brier는 낮을수록 좋다 · RF 튜닝 단계의 Brier 값은 보고서에 기록되지 않았다", size=16)

    # ========================= 54 부록 보고서 Judge 상세 =========================
    sl, top = slide(prs, p(), chapter=APX, headline="보고서 유형별 LLM Judge 점수 상세",
                    source="출처: docs/eval/01_보고서작성에이전트 (Judge gpt-5.6-luna, 2026-09-08)")
    rows = [["보고서 유형", "개선 전", "개선 후"],
            ["미팅 보고서", "97.74", "96.57"],
            ["일일 업무보고", "86.88", "95.16"],
            ["주간 업무보고", "72.50", "87.08"],
            ["월간 업무보고", "66.25", "76.25"]]
    table(sl, M, top, CW, 3.9, rows, col_w=[5.0, 3.546, 3.547], highlight_col=2)
    note(sl, M, top + 4.08, CW,
         "100점 만점 · 공통 43건 · 생성 모델과 Judge 모델이 동일하며 시나리오는 1개, 월간 보고서는 1건이다",
         size=16)

    # ========================= 55 부록 OCR 상세 =========================
    sl, top = slide(prs, p(), chapter=APX, headline="OCR 고도화 전후 문서별 상세",
                    source="출처: OCR 고도화 전후 최종 결과보고서 (2026-09-17 RunPod 재시험)")
    rows = [["대상 (각 50건)", "1차 테스트", "최종 테스트", "변화"],
            ["견적서", "200 / 200 = 100.0%", "200 / 200 = 100.0%", "0.0%p"],
            ["계약서", "50 / 200 = 25.0%", "197 / 200 = 98.5%", "+73.5%p"],
            ["발주서", "502 / 750 = 66.9%", "739 / 750 = 98.5%", "+31.6%p"],
            ["상품설명서", "42 / 50 = 84.0%", "43 / 50 = 86.0%", "+2.0%p"]]
    table(sl, M, top, CW, 3.4, rows, col_w=[2.9, 3.3, 3.3, 2.593], highlight_col=2, size=17)
    note(sl, M, top + 3.6, CW,
         "핵심 구조화 문서(견적서 · 계약서 · 발주서) 합산 — 752 / 1,150 = 65.4% → 1,136 / 1,150 = 98.8% (+33.4%p)",
         size=16, color=SUB)
    note(sl, M, top + 3.98, CW,
         "정규화한 정답값이 OCR 본문에 포함되면 일치로 판정 · 웜 스타트 처리 시간은 명함 2.24초 · PDF 1.23초 (각 5회, 별도 측정)",
         size=15)

    # ========================= 56 부록 출처 일람 =========================
    sl, top = slide(prs, p(), chapter=APX, headline="본문에 쓰인 근거 자료")
    w2, xs2 = grid(2, 0.6)
    srcs = [["프로젝트 기획서 (2026-08)",
             "Grand View Research (2026)",
             "KISDI 2026 · 원자료 Gartner 2025",
             "환율 2026-09-15 기준 1,347.8원",
             "각 사 공식 제품 페이지 (2026-09-15 확인)",
             "SalesLuv 구현 코드 (develop, 2026-09-17)"],
            ["미팅 내용분석 구조 보고서 (2026-08-31 골든셋)",
             "머신러닝 · 딥러닝 학습결과서 (30회 평균)",
             "docs/eval/01_보고서작성에이전트 (2026-09-08)",
             "docs/document-summary-evaluation.md",
             "output/evals/document-summary-rageval",
             "OCR 고도화 전후 최종 결과보고서 (2026-09-17)"]]
    for ci, col in enumerate(srcs):
        bullets(sl, xs2[ci], top + 0.3, w2, col, gap=0.66, size=17)
    note(sl, M, BOT - 0.4, CW,
         "평가용 데이터는 모두 합성 데이터이며 실제 고객 · 영업 데이터를 사용하지 않았다", size=16)

    prs.save(OUT)
    print("saved", OUT, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build()
