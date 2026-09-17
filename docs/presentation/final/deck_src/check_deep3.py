"""SalesLuv_보고서에이전트_Deep_Agents_개선_3장.pptx 검수 — 수치 출처 · 페이지 번호 · 레이아웃.

3장에 적힌 모든 숫자가 근거 덱(Deep Agents 개선 최종본)에 실제로 있는지 대조한다.
없는 숫자가 하나라도 있으면 실패한다 — "수치를 지어내지 않았다"의 기계적 증명이다.
레이아웃은 Pretendard 실측 폭으로 한 줄 넘침을, 상자 겹침으로 카드 충돌을 잡는다.
"""

import math
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from base_v3 import BOT, CW, FOOT_Y, H1_Y, M, SW, TOP
from build_deep3 import AT, OUT
from textwidth import width

SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "eval",
                   "SalesLUV_보고서에이전트_Deep_Agents_개선_최종본.pptx")
CHAPTER = "03 검증과 확장"
PAGES = {1: AT + 1, 2: AT + 2, 3: AT + 3}
TITLES = {1: "토큰 부담과 기간 보고서 완성도가 문제였다",
          2: "감독 · 작성 · 검토로 역할을 나눴다",
          3: "토큰은 절반 이하로, 종합 품질은 +1.94점"}
NUM = re.compile(r"\d[\d,.]*")


def texts(path):
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif getattr(sh, "has_chart", False):
                ch = sh.chart
                out.extend(str(c) for c in ch.plots[0].categories)
                for ser in ch.plots[0].series:
                    for pt in ser._element.iter():
                        if pt.tag.endswith("}v") and pt.text:
                            out.append(pt.text)
            elif sh.has_table:
                out.extend(c.text for r in sh.table.rows for c in r.cells)
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)
    for s in Presentation(path).slides:
        walk(s.shapes)
    return "\n".join(out)


def canon(tok):
    t = tok.replace(",", "")
    if "." in t:
        t = t.rstrip("0").rstrip(".")
    return t


def numbers(blob):
    got = set()
    for tok in NUM.findall(blob):
        got.add(canon(tok))
        if "." in tok:
            for nd in (2, 3, 4):
                try:
                    got.add(canon("%.*f" % (nd, float(tok.replace(",", "")))))
                except ValueError:
                    pass
    return got


def slide_text(sl):
    return "\n".join(sh.text_frame.text for sh in sl.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip())


def boxes(sl):
    return [(sh, sh.left / 914400, sh.top / 914400,
             (sh.width or 0) / 914400, (sh.height or 0) / 914400) for sh in sl.shapes]


def main():
    prs = Presentation(OUT)
    slides = list(prs.slides)
    fail = []

    # ---------- 1. 슬라이드 수 · 화면 비율 ----------
    assert len(slides) == 3, "슬라이드 %d 장 (3 이어야 함)" % len(slides)
    assert (round(prs.slide_width / 914400, 3), round(prs.slide_height / 914400, 3)) == (13.333, 7.5)

    # ---------- 2. 페이지 번호 (v3 뒤 57 · 58 · 59) ----------
    for n, sl in enumerate(slides, 1):
        found = None
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.top is None:
                continue
            if abs(sh.left - Inches(SW - M - 1.0)) < Inches(0.05) and \
               abs(sh.top - Inches(FOOT_Y - 0.03)) < Inches(0.05):
                found = sh.text_frame.text.strip()
        if found != "%02d" % PAGES[n]:
            fail.append("%d 장 페이지 번호 = %r (%02d 이어야 함)" % (n, found, PAGES[n]))

    # ---------- 3. 제목 위치 · 챕터 라벨 (v3 와 같은 좌표) ----------
    for n, sl in enumerate(slides, 1):
        head = [sh.text_frame.text.strip() for sh, x, y, w, h in boxes(sl)
                if sh.has_text_frame and abs(y - H1_Y) < 0.02]
        if head != [TITLES[n]]:
            fail.append("%d 장 제목 = %r" % (n, head))
        if CHAPTER not in slide_text(sl):
            fail.append("%d 장 챕터 라벨 없음" % n)

    # ---------- 4. 수치 대조 — 모든 숫자가 근거 덱에 있어야 한다 ----------
    allowed = numbers(texts(SRC))
    for n, sl in enumerate(slides, 1):
        for tok in NUM.findall(slide_text(sl)):
            c = canon(tok)
            if c in allowed:
                continue
            if c == str(PAGES[n]):                 # 페이지 번호
                continue
            if re.fullmatch(r"0\d", tok):          # 01~04 = 단계 · 챕터 장식 번호
                continue
            fail.append("%d 장의 %r 을 근거 덱에서 찾지 못함" % (n, tok))

    # ---------- 5. 레이아웃 ----------
    for n, sl in enumerate(slides, 1):
        for sh, x, y, w, h in boxes(sl):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:      # 로고
                continue
            if y > BOT:                                      # 출처 · 페이지 번호
                continue
            if x < M - 0.01 or x + w > M + CW + 0.01:
                fail.append("%d 장 좌우 여백 벗어남: %.2f~%.2f" % (n, x, x + w))
            if y < TOP - 0.01 and sh.has_text_frame and sh.text_frame.text.strip() \
               and abs(y - H1_Y) > 0.02 and y > 0.8:
                fail.append("%d 장 본문이 제목 영역 침범: %.2f" % (n, y))
            if y + h > BOT + 0.01 and sh.has_text_frame and sh.text_frame.text.strip():
                fail.append("%d 장 본문 하단 넘침: %.2f" % (n, y + h))

        # 줄바꿈까지 계산한 실제 글 높이가 칸을 넘지 않는가
        # (한 줄 가정으로 잡은 칸에 글이 접히면 카드 밖으로 샌다 — 렌더 없이 잡는 유일한 방법)
        for sh, x, y, w, h in boxes(sl):
            if not (sh.has_text_frame and sh.text_frame.text.strip()) or not sh.width:
                continue
            if y > BOT:
                continue
            tf = sh.text_frame
            inner = w - (tf.margin_left + tf.margin_right) / 914400
            need = 0.0
            wrapped = []
            for para in tf.paragraphs:
                runs = para.runs
                if not runs:
                    continue
                txt = "".join(r.text for r in runs)
                f0 = runs[0].font
                pt = f0.size.pt if f0.size else 15
                wt = 700 if f0.bold else (600 if (f0.name or "").endswith("SemiBold") else
                                          (500 if (f0.name or "").endswith("Medium") else 400))
                nl = max(1, math.ceil(width(txt, pt, wt) / inner - 0.005))
                if nl > 1:
                    wrapped.append(txt[:20])
                need += (para.space_before.pt if para.space_before else 0) / 72
                need += nl * pt * (para.line_spacing or 1.3) / 72
            if need > h + 0.02:
                fail.append("%d 장 글 높이 넘침: %r (%.2f > %.2f in%s)"
                            % (n, sh.text_frame.text.strip()[:22], need, h,
                               " · 줄바꿈 " + " / ".join(wrapped) if wrapped else ""))

        # 글자끼리 겹치지 않는가
        tb = [(sh.text_frame.text.strip()[:18], x, y, w, h) for sh, x, y, w, h in boxes(sl)
              if sh.has_text_frame and sh.text_frame.text.strip() and sh.width]
        for i, (t1, x1, y1, w1, h1) in enumerate(tb):
            for t2, x2, y2, w2, h2 in tb[i + 1:]:
                ox = min(x1 + w1, x2 + w2) - max(x1, x2)
                oy = min(y1 + h1, y2 + h2) - max(y1, y2)
                if ox > 0 and oy > 0 and ox * oy > 0.3 * min(w1 * h1, w2 * h2):
                    fail.append("%d 장 글자 겹침: %r ↔ %r" % (n, t1, t2))

    if fail:
        print("실패 %d 건" % len(fail))
        for f in fail:
            print("  -", f)
        sys.exit(1)
    print("통과 — 3장 · 제목 · 페이지 번호 %d·%d·%d · 수치 출처 · 레이아웃 이상 없음"
          % (PAGES[1], PAGES[2], PAGES[3]))


if __name__ == "__main__":
    main()
