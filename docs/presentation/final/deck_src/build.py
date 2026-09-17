"""SalesLuv 최종 발표 덱 v2 — 키워드 중심 리디자인."""

import os
import sys
import urllib.request

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import cairosvg
from PIL import Image
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from base import *  # noqa: F401,F403
from base import (
    BAR, BLUE, BLUE_500, BLUE_600, BLUE_DEEP, BLUE_LIGHT, BLUE_PALE, BODY_B, BODY_Y, BORDER,
    COL, CW, DARK, DIVIDER, HEADLINE_Y, INK, LOGO, LOGO_WHITE, M, MUTED, RGBColor, SLIDE_H,
    SLIDE_W, SUB, SURFACE, WHITE, add_chart, arrow, blank, card, circle, content_slide, divider,
    frame, gradient_card, line, new_deck, para, pill, textbox, write,
)

OUT = "/Volumes/Jexists/skn30/SKN30-FINAL-1Team-dev/docs/presentation/final/SalesLuv_최종발표_260918.pptx"
HERE = os.path.dirname(os.path.abspath(__file__))
ICON_DIR = os.path.join(HERE, "icons")

CH1 = "01 문제 정의"
CH2 = "02 설계와 구현"
CH3 = "03 검증과 확장"
PH = "[입력 예정]"

TINT = RGBColor(0xF4, 0xF7, 0xFE)
GRAY = RGBColor(0xF7, 0xF8, 0xFA)
L = PP_ALIGN.LEFT
C = PP_ALIGN.CENTER
R = PP_ALIGN.RIGHT


# ---------- 에셋 ----------
def make_white_logo():
    img = Image.open(LOGO).convert("RGBA")
    img.putdata([(255, 255, 255, a) for (_, _, _, a) in img.getdata()])
    img.save(LOGO_WHITE)


BRAND = {"react": "149ECA", "typescript": "3178C6", "vite": "646CFF", "fastapi": "009688",
         "python": "3776AB", "sqlalchemy": "D71F00", "supabase": "1FA97A", "postgresql": "4169E1",
         "openai": "412991", "langchain": "1C3C3C", "scikitlearn": "F7931E", "docker": "2496ED",
         "githubactions": "2088FF", "amazonwebservices": "E8871A", "amazons3": "569A31",
         "amazonec2": "E8871A"}
ICON_SRC = "https://cdn.jsdelivr.net/npm/simple-icons@13/icons/{}.svg"


def icon(slug):
    """simple-icons 공식 로고를 브랜드 컬러 PNG로 받아 캐시."""
    os.makedirs(ICON_DIR, exist_ok=True)
    path = os.path.join(ICON_DIR, slug + ".png")
    if not os.path.exists(path):
        req = urllib.request.Request(ICON_SRC.format(slug), headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=20) as r:
            svg = r.read().decode()
        svg = svg.replace("<svg ", '<svg fill="#%s" ' % BRAND.get(slug, "45515E"), 1)
        cairosvg.svg2png(bytestring=svg.encode(), write_to=path, output_width=320, output_height=320)
    return path


def put_icon(sl, slug, x, y, size):
    return sl.shapes.add_picture(icon(slug), Inches(x), Inches(y), Inches(size), Inches(size))


# ---------- 레이아웃 프리미티브 ----------
def eyebrow(sl, x, y, text, *, color=BLUE, w=4.0):
    card(sl, x, y + 0.045, 0.085, 0.16, fill=color, radius=0.04, shadow=False)
    write(sl, x + 0.2, y, w, 0.26, [(text, dict(size=11, weight=600, color=color, line=1.3))])


def stat(sl, x, y, w, value, label, *, sub=None, size=36, color=INK, align=L):
    write(sl, x, y, w, 0.62, [(value, dict(size=size, weight=700, color=color, line=1.1))], align=align)
    yy = y + size / 72 * 1.12 + 0.06
    lines = [(label, dict(size=11.5, weight=600, color=INK, line=1.3))]
    if sub:
        lines.append((sub, dict(size=9.5, weight=400, color=MUTED, line=1.35, spacing_before=2)))
    write(sl, x, yy, w, 0.6, lines, align=align)


def step(sl, x, y, w, n, title, kw, *, active=False, d=0.44, title_size=14):
    circle(sl, x, y, d, n, fill=(BLUE if active else RGBColor(0xE8, 0xEE, 0xFB)),
           color=(WHITE if active else BLUE_600), size=12.5, weight=700)
    ty = y + d + 0.16
    write(sl, x, ty, w, 0.3,
          [(title, dict(size=title_size, weight=700, color=(BLUE if active else INK), line=1.25))])
    if kw:
        write(sl, x, ty + 0.32, w, 0.9, [(kw, dict(size=11, weight=400, color=SUB, line=1.45))])


def note(sl, x, y, w, text, *, size=11, color=SUB, weight=400):
    return write(sl, x, y, w, 0.5, [(text, dict(size=size, weight=weight, color=color, line=1.5))])


def so_what(sl, y, text, *, x=M, w=CW):
    card(sl, x, y + 0.02, 0.085, 0.26, fill=BLUE, radius=0.04, shadow=False)
    write(sl, x + 0.22, y, w - 0.22, 0.3, [(text, dict(size=12.5, weight=600, color=INK, line=1.35))])


def placeholder(sl, x, y, w, h, text, *, size=11):
    sh = card(sl, x, y, w, h, fill=GRAY, shadow=False, line=BORDER)
    from pptx.oxml import parse_xml
    sh.line._get_or_add_ln().append(parse_xml(
        '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'))
    tf = sh.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, weight=500, color=MUTED, line=1.4, align=C, first=True)
    return sh


def grid(n, total_w, gap):
    w = (total_w - gap * (n - 1)) / n
    return w, [M + i * (w + gap) for i in range(n)]


def build():
    make_white_logo()
    prs = new_deck()
    p = 1

    # ================= 01 표지 =================
    sl = blank(prs)
    sl.shapes.add_picture(LOGO, Inches(SLIDE_W - M - 1.03), Inches(0.44), Inches(1.03), Inches(0.24))
    eyebrow(sl, 0.9, 1.95, "SKN30 1Team · CassTerra")
    write(sl, 0.9, 2.3, 7.0, 1.0, [("SalesLuv", dict(size=62, weight=700, color=INK, line=1.05))])
    write(sl, 0.9, 3.45, 6.2, 0.4,
          [("영업의 흐름을 연결하는 멀티에이전트 CRM", dict(size=20, weight=500, color=SUB, line=1.35))])
    line(sl, 0.9, 4.15, 4.6, 4.15, color=BORDER)
    write(sl, 0.9, 4.35, 6.0, 0.8,
          [("2026.09.18 (FRI) 15:00", dict(size=13, weight=600, color=INK, line=1.5)),
           ("천성배 · 정주애 · 박제섭 · 박지유", dict(size=13, weight=400, color=MUTED, line=1.5, spacing_before=3))])

    gc = gradient_card(sl, 7.55, 1.3, 5.28, 4.9, radius=0.3)
    sl.shapes.add_picture(LOGO_WHITE, Inches(7.55 + (5.28 - 2.0) / 2), Inches(2.55), Inches(2.0), Inches(0.47))
    write(sl, 7.95, 3.5, 4.48, 1.0,
          [("CONNECT THE HISTORY.", dict(size=17, weight=700, color=WHITE, line=1.4)),
           ("APPROVE THE NEXT MOVE.", dict(size=17, weight=700, color=WHITE, line=1.4))], align=C)
    write(sl, 7.95, 4.75, 4.48, 0.4,
          [("기록을 잇고, 승인된 다음 행동으로", dict(size=12, weight=400, color=BLUE_PALE, line=1.4))], align=C)

    # ================= 02 목차 =================
    p += 1
    sl = content_slide(prs, "목차", "문제 → 구현 → 검증 순서로 설명합니다", "15분, 세 파트", p)
    w3, xs = grid(3, CW, 0.5)
    toc = [("01", "문제 정의", "팀 · 주제 · 문제 · 시장 · 해결책", "왜 이 문제를 골랐는가"),
           ("02", "설계와 구현", "시나리오 · 시연 · 스택 · 아키텍처 · 에이전트", "무엇을 어떻게 만들었는가"),
           ("03", "검증과 확장", "Agent 평가 · 실패와 개선 · 향후 계획 · 결론", "제대로 동작하는지 어떻게 확인했는가")]
    for i, (n, t, kw, lead) in enumerate(toc):
        x = xs[i]
        if i:
            line(sl, x - 0.25, BODY_Y + 0.25, x - 0.25, BODY_B - 0.25, color=BORDER)
        write(sl, x, BODY_Y + 0.5, w3, 1.1, [(n, dict(size=64, weight=700, color=BLUE_PALE, line=1.05))])
        write(sl, x, BODY_Y + 1.75, w3, 0.4, [(t, dict(size=24, weight=700, color=INK, line=1.3))])
        write(sl, x, BODY_Y + 2.3, w3, 0.4, [(lead, dict(size=13, weight=500, color=BLUE_600, line=1.4))])
        line(sl, x, BODY_Y + 2.85, x + w3 - 0.4, BODY_Y + 2.85, color=DIVIDER)
        write(sl, x, BODY_Y + 3.05, w3 - 0.3, 0.9, [(kw, dict(size=12, weight=400, color=SUB, line=1.6))])

    # ================= D1 =================
    p += 1
    divider(prs, "PART 01", "문제 정의", "왜 이 프로젝트가 필요한가", p)

    # ================= 03 팀 =================
    p += 1
    sl = content_slide(prs, CH1, "네 명이 기획부터 배포까지 나눠 맡았다", "PM · 풀스택 · 에이전트 · 시연", p)
    w4, xs = grid(4, CW, 0.32)
    team = [("천성배", ["P.M", "OCR", "자료요약 Agent"]),
            ("정주애", ["Full Stack", "Infra"]),
            ("박제섭", ["기술리더", "미팅분석 Agent", "업무보고 Agent"]),
            ("박지유", ["영업계약 Agent", "일정관리 Agent", "시연영상"])]
    ph_h = 2.1
    for i, (name, roles) in enumerate(team):
        x = xs[i]
        placeholder(sl, x, BODY_Y + 0.1, w4, ph_h, "PHOTO", size=12)
        write(sl, x, BODY_Y + ph_h + 0.32, w4, 0.35,
              [(name, dict(size=18, weight=700, color=INK, line=1.25))])
        cy = BODY_Y + ph_h + 0.75
        for r_i, role in enumerate(roles):
            pill(sl, x, cy + r_i * 0.36, min(w4, 0.11 * len(role) + 0.62), 0.28, role,
                 fill=(TINT if r_i == 0 else SURFACE), color=(BLUE_600 if r_i == 0 else SUB), size=10)

    # ================= 04 주제 =================
    p += 1
    sl = content_slide(prs, CH1, "담당자의 고객 맥락을 팀이 이어받게 한다",
                       "미팅 → 계약까지 잇는 멀티에이전트 CRM", p,
                       source="출처: 프로젝트 기획서(2026-08)")
    card(sl, M, BODY_Y + 0.15, CW, 1.15, fill=TINT, shadow=False)
    write(sl, M + 0.5, BODY_Y + 0.52, CW - 1.0, 0.6,
          [("영업의 시작부터 끝까지, 한 번에 해결하는 멀티에이전트 CRM",
            dict(size=24, weight=700, color=BLUE_DEEP, line=1.35))], align=C)
    w3, xs = grid(3, CW, 0.3)
    defs = [("국내 B2B 영업", "고객·딜·일정·보고가 흩어진 조직"),
            ("멀티에이전트", "판단·초안·제안을 다섯으로 분담"),
            ("Human-in-the-loop", "AI는 초안까지, 확정은 사람의 승인으로")]
    for i, (t, d) in enumerate(defs):
        x = xs[i]
        line(sl, x, BODY_Y + 1.95, x + w3, BODY_Y + 1.95, color=BLUE_PALE, width=2)
        write(sl, x, BODY_Y + 2.15, w3, 0.35, [(t, dict(size=16, weight=700, color=INK, line=1.3))])
        write(sl, x, BODY_Y + 2.6, w3, 0.6, [(d, dict(size=11.5, weight=400, color=SUB, line=1.5))])
    note(sl, M, BODY_B - 0.52, CW,
         "CRM — 고객과 나눈 이야기를 기억하는 시스템    ·    딜 — 하나의 거래 기회",
         size=10.5, color=MUTED)

    # ================= 05 문제 =================
    p += 1
    sl = content_slide(prs, CH1, "영업 맥락은 담당자가 바뀌는 순간 끊긴다",
                       "정보 분산 → 인수인계 맥락 소실", p,
                       source="출처: 프로젝트 기획서(2026-08) · 영업사원 출신 PM의 현장 경험")
    w5, xs = grid(5, CW, 0.28)
    probs = [("01", "영업정보 분산", "고객·계약·매출이 따로"),
             ("02", "영업흐름 단절", "리드→계약→매출이 끊김"),
             ("03", "판단 공백", "진행·예측이 안 보임"),
             ("04", "인수인계 맥락 소실", "관계와 노하우가 사라짐"),
             ("05", "범용 CRM 한계", "B2B 영업엔 부족")]
    ny = BODY_Y + 0.35
    line(sl, xs[0] + 0.22, ny + 0.22, xs[4] + 0.22, ny + 0.22, color=BORDER, width=1.5)
    for i, (n, t, d) in enumerate(probs):
        step(sl, xs[i], ny, w5, n, t, d, active=(i == 3), title_size=13.5)
    card(sl, M, BODY_B - 1.55, CW, 1.05, fill=TINT, shadow=False)
    card(sl, M, BODY_B - 1.55, 0.09, 1.05, fill=BLUE, radius=0.045, shadow=False)
    write(sl, M + 0.45, BODY_B - 1.32, CW - 0.9, 0.7,
          [("고객 맥락과 영업 노하우는 담당자 개인의 암묵지로만 쌓인다",
            dict(size=17, weight=700, color=INK, line=1.35)),
           ("조직 자산으로 남지 않아 담당자가 바뀌면 함께 사라진다",
            dict(size=12, weight=400, color=SUB, line=1.4, spacing_before=5))])

    # ================= 06 시장 =================
    p += 1
    sl = content_slide(prs, CH1, "시장은 성장, 진입 조건은 맥락과 근거의 연결",
                       "영업 CRM 단독 통계 없음 → 성장성·진입 조건으로 판단", p,
                       source="출처: Grand View Research(2026), KISDI 2026(원자료 Gartner 2025) · 환율 1,347.8원")
    cw_ = 6.6
    eyebrow(sl, M, BODY_Y + 0.05, "국내 CRM 시장 규모 (조원)")
    ch = add_chart(sl, COL, M - 0.15, BODY_Y + 0.4, cw_, 2.8, ["2025", "2026", "2033"],
                   [("국내 CRM 시장", [2.51, 2.75, 5.46])], label_fmt="0.00", gap=110, y_max=6.0)
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    stat(sl, rx, BODY_Y + 0.15, rw / 2 - 0.1, "10.3%", "CRM CAGR", sub="2026–2033", color=BLUE, size=34)
    stat(sl, rx + rw / 2 + 0.1, BODY_Y + 0.15, rw / 2 - 0.1, "17.1%", "SaaS CAGR", sub="2025–2029", color=BLUE_500, size=34)
    line(sl, rx, BODY_Y + 1.4, rx + rw, BODY_Y + 1.4, color=BORDER)
    conds = [("경쟁 현실", "AI 미팅 요약·후속조치는 이미 존재"),
             ("초기 ICP", "문서·견적·계약이 복잡한 B2B 영업팀"),
             ("차별화 조건", "근거 검색 + 데이터 연결 + 담당자 승인")]
    for i, (t, d) in enumerate(conds):
        y = BODY_Y + 1.6 + i * 0.85
        write(sl, rx, y, rw, 0.28, [(t, dict(size=13, weight=700, color=BLUE_600, line=1.3))])
        write(sl, rx, y + 0.3, rw, 0.45, [(d, dict(size=11.5, weight=400, color=SUB, line=1.45))])
    note(sl, M, BODY_B - 0.5, cw_, "CRM·SaaS는 정의가 달라 합산하지 않음",
         size=10, color=MUTED)

    # ================= 07 비교 =================
    p += 1
    sl = content_slide(prs, CH1, "차이는 AI 유무가 아니라 승인이 순환하는 구조다",
                       "데이터 관리에서 다음 행동으로", p,
                       source="출처: 각 사 공식 제품 페이지(2026-09-15 확인) · 기능군 판정은 팀 조사 기준")
    tw = 7.9
    rows = [("기능군", "세일즈맵", "핑거세일즈", "세일즈인사이트", "SalesLuv"),
            ("고객·거래처 관리", "●", "●", "●", "●"),
            ("파이프라인·딜 관리", "●", "●", "●", "●"),
            ("영업 실적·매출 분석", "●", "●", "△", "●"),
            ("AI 기능", "△", "●", "–", "●"),
            ("사용자 승인 후 업무 반영", "–", "●", "–", "●"),
            ("승인 결과 → 다음 행동 순환", "–", "–", "–", "●")]
    th = 3.3
    tbl = sl.shapes.add_table(len(rows), 5, Inches(M), Inches(BODY_Y + 0.1), Inches(tw), Inches(th)).table
    tbl.columns[0].width = Inches(3.1)
    for i in range(1, 5):
        tbl.columns[i].width = Inches((tw - 3.1) / 4)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(th / len(rows))
        for ci, val in enumerate(row):
            cell = tbl.cell(r, ci)
            cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            last = r == len(rows) - 1
            para(cell.text_frame, val, size=11.5 if r else 11,
                 weight=700 if (last and ci) or r == 0 else 400,
                 color=(BLUE if ci == 4 and r else INK if r == 0 or ci == 0 else SUB),
                 line=1.2, first=True, align=L if ci == 0 else C)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DIVIDER if r == 0 else (TINT if last else WHITE)
    note(sl, M, BODY_Y + th + 0.22, tw, "● 제공 · △ 일부 · – 공개 자료에서 확인되지 않음", size=9.5, color=MUTED)
    rx = M + tw + 0.4
    rw = CW - tw - 0.4
    eyebrow(sl, rx, BODY_Y + 0.1, "SalesLuv만의 구조", w=rw)
    diffs = ["근거 기반 추천", "Human-in-the-loop", "다음 행동으로 순환"]
    subs = ["보고서·자료 검색 + 출처", "AI 제안 → 담당자 승인", "승인이 다음 추천을 부른다"]
    for i in range(3):
        y = BODY_Y + 0.65 + i * 1.1
        card(sl, rx, y, rw, 0.9, fill=WHITE, line=BORDER)
        card(sl, rx, y, 0.08, 0.9, fill=[BLUE_PALE, BLUE_LIGHT, BLUE][i], radius=0.04, shadow=False)
        write(sl, rx + 0.28, y + 0.18, rw - 0.5, 0.28, [(diffs[i], dict(size=13.5, weight=700, color=INK, line=1.25))])
        write(sl, rx + 0.28, y + 0.5, rw - 0.5, 0.35, [(subs[i], dict(size=10.5, weight=400, color=SUB, line=1.35))])

    # ================= 08 해결책 =================
    p += 1
    sl = content_slide(prs, CH1, "기록이 초안이 되고, 승인이 다음 일정이 된다",
                       "딜 중심 — 기록 → 판단 → 일정 → 문서", p,
                       source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    w4, xs = grid(4, CW, 0.5)
    sols = [("01", "미팅 기록 입력", "원문 · 음성 STT · 이미지 OCR"),
            ("02", "AI 초안 + 딜 승산", "딜별 귀속 · 13개 특성 · 성사 확률"),
            ("03", "사람이 확정", "근거 대조 · 확정 · 팀장 검토"),
            ("04", "다음 미팅 · 브리핑", "날짜 추천 · 캘린더 승인")]
    cy = BODY_Y + 0.7
    ch_h = 1.5
    for i, (n, t, d) in enumerate(sols):
        x = xs[i]
        card(sl, x, cy, w4, ch_h, fill=WHITE, line=BORDER)
        circle(sl, x + 0.28, cy - 0.22, 0.44, n, fill=BLUE, size=12.5)
        write(sl, x + 0.28, cy + 0.42, w4 - 0.56, 0.3, [(t, dict(size=15, weight=700, color=INK, line=1.3))])
        write(sl, x + 0.28, cy + 0.82, w4 - 0.56, 0.8, [(d, dict(size=11.5, weight=400, color=SUB, line=1.5))])
        if i < 3:
            arrow(sl, x + w4 + 0.1, cy + ch_h / 2, x + w4 + 0.4, cy + ch_h / 2, color=BLUE_LIGHT)
    ry = cy + ch_h + 0.45
    line(sl, M + w4 / 2, ry, M + w4 / 2, ry + 0.45, color=BLUE_PALE, width=1.5, dashed=True)
    line(sl, M + w4 / 2, ry + 0.45, xs[3] + w4 / 2, ry + 0.45, color=BLUE_PALE, width=1.5, dashed=True)
    arrow(sl, xs[3] + w4 / 2, ry + 0.45, xs[3] + w4 / 2, ry + 0.02, color=BLUE_PALE)
    write(sl, M + w4 / 2 + 0.3, ry + 0.52, 8.0, 0.3,
          [("승인 결과가 다시 01로 — 모든 기록은 딜 하나로 연결",
            dict(size=12, weight=600, color=BLUE_600, line=1.35))])

    # ================= 09 기대효과 =================
    p += 1
    sl = content_slide(prs, CH1, "반복 작성은 줄이고, 맥락은 조직에 남긴다",
                       "정량 효과는 파일럿 측정 후 기입", p,
                       source="출처: 프로젝트 기획서 9장(파일럿 KPI 정의 · 목표는 검증할 가설, 실측값 아님)")
    lw = 7.4
    eyebrow(sl, M, BODY_Y + 0.05, "파일럿 측정 지표")
    kw_, kxs = grid(2, lw, 0.3)
    kpis = [("보고서 작성시간", "활성 작업시간 중앙값"),
            ("당일 제출률", "당일 제출 ÷ 보고 대상"),
            ("필수정보 초회 완성률", "첫 제출 필수항목 충족"),
            ("팀장 검토시간", "검토 작업시간 중앙값")]
    for i, (t, d) in enumerate(kpis):
        x = kxs[i % 2]
        y = BODY_Y + 0.45 + (i // 2) * 1.25
        placeholder(sl, x, y, kw_, 1.05, "")
        write(sl, x + 0.22, y + 0.15, kw_ - 0.44, 0.35, [("[수치 입력 예정]", dict(size=15, weight=700, color=MUTED, line=1.25))])
        write(sl, x + 0.22, y + 0.53, kw_ - 0.44, 0.45,
              [(t, dict(size=12, weight=600, color=INK, line=1.3)),
               (d, dict(size=9.5, weight=400, color=MUTED, line=1.3, spacing_before=2))])
    chips = ["자동 집계·검증", "승인 기반 CRM·캘린더 반영", "영업 암묵지의 조직 자산화"]
    cx = M
    for t in chips:
        w_ = 0.105 * len(t) + 0.55
        pill(sl, cx, BODY_B - 0.52, w_, 0.32, t, fill=TINT, color=BLUE_600, size=11)
        cx += w_ + 0.15
    rx = M + lw + 0.5
    rw = CW - lw - 0.5
    eyebrow(sl, rx, BODY_Y + 0.05, "반복 작성을 줄이는 보고서 계층", w=rw)
    levels = [("미팅 보고서", BLUE_PALE, INK), ("일일 업무보고", BLUE_LIGHT, WHITE),
              ("주간 업무보고", BLUE_500, WHITE), ("월간 업무보고", BLUE, WHITE)]
    for i, (t, fc, tc) in enumerate(levels):
        y = BODY_Y + 0.5 + i * 0.72
        ins = i * 0.2
        sh = card(sl, rx + ins, y, rw - ins * 2, 0.56, fill=fc, shadow=False)
        tf = sh.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(tf, t, size=12.5, weight=600, color=tc, line=1.2, align=C, first=True)
        if i < 3:
            arrow(sl, rx + rw / 2, y + 0.58, rx + rw / 2, y + 0.7, color=BLUE_PALE, width=1.2)
    note(sl, rx, BODY_Y + 3.5, rw, "각 단계는 아래 단계의 확정본을 원천으로 사용", size=10.5, color=MUTED)

    # ================= D2 =================
    p += 1
    divider(prs, "PART 02", "설계와 구현", "그래서 어떻게 해결했는가", p)

    # ================= 10 사용자 시나리오 =================
    p += 1
    sl = content_slide(prs, CH2, "AI는 준비·초안·제안까지, 확정은 사람이 한다",
                       "미팅 기록 → 다음 미팅 브리핑", p,
                       source="출처: SalesLuv 구현 코드(develop, 2026-09-17)")
    flow = [("01", "미팅 준비", "계약관리 추천 → 캘린더 승인"),
            ("02", "사전 브리핑", "과거 보고서·자료 RAG"),
            ("03", "미팅 후 입력", "원문 · STT · OCR"),
            ("04", "보고서 확정", "사람이 확정해야 저장"),
            ("05", "보고서 누적", "RAG 적재 · 추천 트리거"),
            ("06", "기간 보고", "일·주·월 → 팀장 검토")]
    w3, xs3 = grid(3, CW, 0.4)
    for i, (n, t, d) in enumerate(flow):
        x = xs3[i % 3]
        y = BODY_Y + 0.45 + (i // 3) * 2.25
        circle(sl, x, y, 0.46, n, fill=(BLUE if i in (3,) else RGBColor(0xE8, 0xEE, 0xFB)),
               color=(WHITE if i in (3,) else BLUE_600), size=12.5)
        write(sl, x + 0.62, y + 0.05, w3 - 0.7, 0.32, [(t, dict(size=15, weight=700, color=INK, line=1.25))])
        write(sl, x + 0.62, y + 0.42, w3 - 0.7, 0.5, [(d, dict(size=11.5, weight=400, color=SUB, line=1.45))])
        if i % 3 < 2:
            arrow(sl, x + w3 + 0.05, y + 0.23, x + w3 + 0.3, y + 0.23, color=BLUE_PALE, width=1.2)
        if i == 2:
            line(sl, M, BODY_Y + 1.95, M + CW, BODY_Y + 1.95, color=DIVIDER)
    so_what(sl, BODY_B - 0.5, "확정 결과가 다음 미팅 준비로 순환")

    # ================= 11 시연 =================
    p += 1
    sl = content_slide(prs, CH2, "한 번의 영업 사이클을 실제 화면으로 따라간다",
                       "[시연 동선 한 줄 입력 예정]", p)
    vw = 8.2
    vh = vw * 9 / 16
    placeholder(sl, M, BODY_Y + 0.1, vw, vh, "시연 영상 (16:9)", size=14)
    rx = M + vw + 0.5
    rw = CW - vw - 0.5
    eyebrow(sl, rx, BODY_Y + 0.15, "시연 체크포인트", w=rw)
    cps = ["자료실 업로드 → 요약", "캘린더 추천 승인", "AI 브리핑 확인", "미팅 원문 → 초안 생성", "보고서 확정 → 다음 추천"]
    for i, t in enumerate(cps):
        y = BODY_Y + 0.65 + i * 0.66
        circle(sl, rx, y, 0.3, str(i + 1), fill=TINT, color=BLUE_600, size=10.5)
        write(sl, rx + 0.42, y + 0.02, rw - 0.42, 0.4, [(t, dict(size=12, weight=500, color=INK, line=1.35))])
    note(sl, rx, BODY_B - 0.45, rw, "시연 순서는 영상 확정 후 조정", size=10, color=MUTED)

    # ================= 12 스택 =================
    p += 1
    sl = content_slide(prs, CH2, "AI 생태계를 중심에 두고 스택을 맞췄다",
                       "Python AI 생태계 기준", p,
                       source="출처: frontend/package.json, backend/uv.lock, 운영 설정값(2026-09-11 확인)")
    groups = [("FRONTEND", [("react", "React 19"), ("typescript", "TypeScript"), ("vite", "Vite")], []),
              ("BACKEND", [("fastapi", "FastAPI"), ("python", "Python 3.13"), ("sqlalchemy", "SQLAlchemy")], []),
              ("DATA · AUTH", [("supabase", "Supabase"), ("postgresql", "PostgreSQL")], ["pgvector", "Auth · Storage"]),
              ("AI · ML", [("openai", "OpenAI API"), ("langchain", "LangChain"), ("scikitlearn", "scikit-learn")],
               ["DeepAgents · CatBoost", "RunPod OCR"]),
              ("DEVOPS", [("docker", "Docker"), ("githubactions", "GitHub Actions"), ("amazonwebservices", "AWS")], [])]
    w5, xs5 = grid(5, CW, 0.22)
    top = BODY_Y + 0.1
    gh = 3.8
    for gi, (label, items, extra) in enumerate(groups):
        x = xs5[gi]
        card(sl, x, top, w5, gh, fill=GRAY, shadow=False)
        write(sl, x + 0.22, top + 0.24, w5 - 0.44, 0.3,
              [(label, dict(size=11, weight=700, color=BLUE_600, line=1.3))])
        line(sl, x + 0.22, top + 0.62, x + w5 - 0.22, top + 0.62, color=BORDER)
        for i, (slug, name) in enumerate(items):
            y = top + 0.82 + i * 0.68
            bx = card(sl, x + 0.22, y, w5 - 0.44, 0.54, fill=WHITE, shadow=False)
            put_icon(sl, slug, x + 0.38, y + 0.11, 0.32)
            write(sl, x + 0.86, y + 0.14, w5 - 1.1, 0.3,
                  [(name, dict(size=11.5, weight=600, color=INK, line=1.25))])
        ey = top + 0.82 + len(items) * 0.68 + 0.06
        for j, t in enumerate(extra):
            write(sl, x + 0.3, ey + j * 0.3, w5 - 0.5, 0.28,
                  [("· " + t, dict(size=10.5, weight=400, color=SUB, line=1.35))])
    note(sl, M, top + gh + 0.25, CW,
         "LLM gpt-5.6-luna · STT gpt-4o-transcribe · OCR PaddleOCR(RunPod) — adapter로 교체 가능",
         size=11, color=SUB)

    # ================= 13 아키텍처 =================
    p += 1
    sl = content_slide(prs, CH2, "입구는 하나, 오래 걸리는 AI 작업은 따로 뗐다",
                       "LLM 생성 = 큐 + 워커 · OCR·STT = API 직접 호출", p,
                       source="출처: deploy/backend/deploy.sh, backend/app/services/agent_worker.py")
    lane_x, lane_w = 3.15, 5.4          # 중앙 흐름
    side_x, side_w = 9.05, 3.78         # 오른쪽 부속
    rail_x = M + 0.05
    rows = [
        # (y, h, kind, title, sub, icon, side)
        (BODY_Y + 0.0, 0.46, "user", "사용자", None, None, None),
        (BODY_Y + 0.62, 0.66, "edge", "CloudFront", "단일 HTTPS 진입점", "amazonwebservices",
         ("amazons3", "S3 — React SPA", "정적 호스팅 · Private")),
        (BODY_Y + 1.46, 0.66, "app", "EC2 · Nginx · FastAPI", "API 서버 · SSE", "amazonec2",
         ("openai", "OCR · STT 동기 호출", "RunPod · gpt-4o-transcribe")),
        (BODY_Y + 2.3, 0.74, "data", "Supabase", "PostgreSQL · pgvector · Auth · Storage · agent_run 큐", "supabase", None),
        (BODY_Y + 3.22, 0.66, "async", "Agent Worker", "큐 폴링 · 별도 컨테이너", "docker",
         ("openai", "OpenAI gpt-5.6-luna", "보고서·미팅분석 생성")),
    ]
    zone_label = {"user": ("CLIENT", MUTED), "edge": ("EDGE", BLUE_600), "app": ("APP", BLUE_600),
                  "data": ("DATA", BLUE), "async": ("ASYNC", BLUE_500)}
    fills = {"user": GRAY, "edge": TINT, "app": TINT, "data": RGBColor(0xE3, 0xEC, 0xFD), "async": TINT}
    prev = None
    for (y, h, kind, title, sub, ic, side) in rows:
        w_ = lane_w if kind != "user" else 2.4
        x_ = lane_x if kind != "user" else lane_x + (lane_w - 2.4) / 2
        box = card(sl, x_, y, w_, h, fill=fills[kind], shadow=False,
                   line=(BLUE_PALE if kind in ("data", "app", "edge") else BORDER))
        tx = x_ + (0.72 if ic else 0.28)
        if ic:
            put_icon(sl, ic, x_ + 0.26, y + (h - 0.3) / 2, 0.3)
        lines = [(title, dict(size=14, weight=700, color=INK, line=1.25))]
        if sub:
            lines.append((sub, dict(size=10.5, weight=400, color=SUB, line=1.35, spacing_before=3)))
        write(sl, tx, y + (0.14 if sub else (h - 0.22) / 2), w_ - (tx - x_) - 0.2, h, lines)
        lb, lc = zone_label[kind]
        write(sl, rail_x, y + (h - 0.2) / 2, 2.3, 0.24, [(lb, dict(size=9.5, weight=700, color=lc, line=1.2))])
        if prev is not None:
            arrow(sl, lane_x + lane_w / 2, prev, lane_x + lane_w / 2, y - 0.02, color=BLUE_LIGHT, width=1.5)
        prev = y + h
        if side:
            s_ic, s_t, s_d = side
            sy = y - 0.02
            card(sl, side_x, sy, side_w, h + 0.04, fill=WHITE, shadow=False, line=BORDER)
            put_icon(sl, s_ic, side_x + 0.24, sy + (h + 0.04 - 0.28) / 2, 0.28)
            write(sl, side_x + 0.66, sy + 0.14, side_w - 0.9, h,
                  [(s_t, dict(size=12, weight=600, color=INK, line=1.25)),
                   (s_d, dict(size=10, weight=400, color=MUTED, line=1.3, spacing_before=3))])
            arrow(sl, lane_x + lane_w + 0.05, y + h / 2, side_x - 0.05, y + h / 2,
                  color=BLUE_PALE, width=1.2, dashed=True)
    note(sl, M, BODY_B - 0.3, CW,
         "배포 — GitHub Actions → OIDC → SSM · 무중단(blue/green)",
         size=10.5, color=MUTED)

    # ================= 14 에이전트 5종 =================
    p += 1
    sl = content_slide(prs, CH2, "다섯 에이전트가 판단·초안·제안을 나눠 맡는다",
                       "제안·초안까지 에이전트 · 확정은 사람", p,
                       source="출처: backend/app/agents/ (develop, 2026-09-17)")
    agents = [("미팅분석", "판단 생성", "미팅 원문", "딜별 귀속 · 13개 특성 · 성사 확률", "확정 시 저장"),
              ("영업·계약관리", "상태 허브", "위험 신호 7종 · 이력", "다음 미팅 제안 · 근거 브리핑", "캘린더 승인"),
              ("보고서 작성", "실행 지원", "원문 · 하위 확정 보고서", "작성 → 검토 → 수정 초안", "확정 · 팀장 검토"),
              ("일정관리", "실행 지원", "저장된 추천 · 딜 상태", "유효 / 재생성 / 무시 판단", "승인 · 거절"),
              ("자료요약", "실행 지원", "PDF · DOCX · 이미지", "구조화 요약 · RAG 청크", "원문 대조")]
    rh = 0.74
    gap = 0.12
    top = BODY_Y + 0.12
    for i, (name, tag, src, out, human) in enumerate(agents):
        y = top + i * (rh + gap)
        card(sl, M, y, CW, rh, fill=WHITE, line=BORDER)
        card(sl, M, y, 0.075, rh, fill=(BLUE if i == 0 else BLUE_LIGHT if i == 1 else BLUE_PALE),
             radius=0.04, shadow=False)
        write(sl, M + 0.3, y + 0.16, 1.75, 0.3, [(name, dict(size=14, weight=700, color=INK, line=1.25))])
        write(sl, M + 0.3, y + 0.47, 1.75, 0.24, [(tag, dict(size=9.5, weight=500, color=MUTED, line=1.2))])
        write(sl, M + 2.2, y + (rh - 0.24) / 2, 2.55, 0.3, [(src, dict(size=11.5, weight=400, color=SUB, line=1.3))])
        arrow(sl, M + 4.85, y + rh / 2, M + 5.25, y + rh / 2, color=BLUE_LIGHT, width=1.2)
        write(sl, M + 5.45, y + (rh - 0.24) / 2, 4.0, 0.3, [(out, dict(size=12, weight=600, color=INK, line=1.3))])
        pill(sl, M + 9.65, y + (rh - 0.3) / 2, 2.6, 0.3, human, fill=TINT, color=BLUE_600, size=10)

    # ================= 15 Agent Flow =================
    p += 1
    sl = content_slide(prs, CH2, "트리거와 작업 큐로 잇고, 사람이 닫는다",
                       "전역 Supervisor 없음 · 코드와 큐로 연결", p,
                       source="출처: backend/app/services/agent_runs.py · agent_worker.py")
    cols = [("업무 트리거", ["자료 업로드", "미팅 보고서 생성", "보고서 확정", "일정 등록", "딜 단계 이동 외 3건"], GRAY),
            ("작업 큐", ["agent_run · PostgreSQL", "부모 · 자식 실행", "최대 2회 · lease 90초"], TINT),
            ("에이전트", ["미팅 내용분석", "보고서 작성 (DeepAgents)", "딜 특성 + ML 예측", "계약관리 → 일정관리", "브리핑 · 자료요약"], WHITE),
            ("사람", ["보고서 확정", "추천 승인 → 일정 등록", "팀장 검토 · 반려"], TINT)]
    w4, xs4 = grid(4, CW, 0.55)
    top = BODY_Y + 0.3
    ch_h = 3.0
    for i, (label, items, fc) in enumerate(cols):
        x = xs4[i]
        card(sl, x, top, w4, ch_h, fill=fc, shadow=(fc is WHITE), line=(BLUE_PALE if i == 2 else None))
        write(sl, x + 0.25, top + 0.26, w4 - 0.5, 0.3,
              [(label, dict(size=13, weight=700, color=(BLUE if i else INK), line=1.25))])
        line(sl, x + 0.25, top + 0.66, x + w4 - 0.25, top + 0.66, color=BORDER)
        for j, t in enumerate(items):
            write(sl, x + 0.25, top + 0.85 + j * 0.46, w4 - 0.5, 0.42,
                  [("· " + t, dict(size=11.5, weight=(600 if i == 2 else 400),
                                   color=(INK if i == 2 else SUB), line=1.4))])
        if i < 3:
            arrow(sl, x + w4 + 0.08, top + ch_h / 2, x + w4 + 0.47, top + ch_h / 2, color=BLUE_LIGHT)
    ry = top + ch_h + 0.3
    arrow(sl, xs4[3] + w4 / 2, ry, xs4[0] + w4 / 2, ry, color=BLUE_PALE, width=1.5, dashed=True)
    write(sl, M + 1.0, ry + 0.08, CW - 2.0, 0.3,
          [("사람의 확정이 다시 트리거가 된다", dict(size=12, weight=600, color=BLUE_600, line=1.3))],
          align=C)

    # ================= D3 =================
    p += 1
    divider(prs, "PART 03", "검증과 확장", "제대로 동작하는지 어떻게 확인했는가", p)

    # ---------- 공통: 지표 스트립 ----------
    def strip(sl, y, items, *, x=M, w=CW, h=1.1, featured=None):
        n = len(items)
        cw_ = w / n
        for i, (v, lb, sub) in enumerate(items):
            cx = x + i * cw_
            if featured == i:
                gradient_card(sl, cx + 0.02, y, cw_ - 0.04, h, radius=0.16)
                write(sl, cx + 0.2, y + 0.16, cw_ - 0.44, 0.5,
                      [(v, dict(size=30, weight=700, color=WHITE, line=1.1))])
                write(sl, cx + 0.2, y + 0.66, cw_ - 0.44, 0.4,
                      [(lb, dict(size=11.5, weight=600, color=WHITE, line=1.25)),
                       (sub, dict(size=9.5, weight=400, color=BLUE_PALE, line=1.25, spacing_before=2))])
            else:
                if i:
                    line(sl, cx, y + 0.12, cx, y + h - 0.12, color=BORDER)
                stat(sl, cx + 0.22, y + 0.1, cw_ - 0.4, v, lb, sub=sub, size=30, color=INK)

    # ================= 16 미팅분석 평가 =================
    p += 1
    sl = content_slide(prs, CH3, "딜 귀속 84.7%, 특성 추출 74.8%가 다음 과제",
                       "딜 승산 모델 AUC +0.037 (30회 평균)", p,
                       source="출처: 미팅 내용분석 구조 보고서(2026-08-31 골든셋), 머신러닝·딥러닝 학습결과서(30회 평균)")
    strip(sl, BODY_Y + 0.05, [("84.7%", "근거·딜 일치", "94 / 111 구간"),
                              ("74.8%", "딜 특성 필드 정확도", "350 / 468 필드"),
                              ("82.6%", "필수 사실 보존", "38 / 46"),
                              ("0 / 36", "13개 특성 전부 정답", "가장 큰 과제")])
    y2 = BODY_Y + 1.45
    cw_ = 8.0
    eyebrow(sl, M, y2, "딜 승산 예측 — 단계별 성능 (30회 평균)")
    add_chart(sl, COL, M - 0.15, y2 + 0.35, cw_, 2.55, ["RF 기준선", "RF 튜닝", "Stacking_LR (최종)"],
              [("Accuracy", [0.6958, 0.7338, 0.7360]), ("ROC-AUC", [0.7409, 0.7756, 0.7783])],
              label_fmt="0.000", gap=80, y_min=0.6, y_max=0.82, colors=(BLUE_PALE, BLUE))
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    eyebrow(sl, rx, y2, "평가 설계와 한계", w=rw)
    for i, t in enumerate(["합성 미팅 12회 · 딜 36개 (2026-08-31)",
                           "AI 의미 검수 — 사람 판정 아님",
                           "공개 B2B 448건 · 30회 반복",
                           "Brier 0.2097 → 0.1880",
                           "영문 데이터 — 한국어 성능은 별도"]):
        write(sl, rx, y2 + 0.45 + i * 0.5, rw, 0.45,
              [("· " + t, dict(size=11, weight=400, color=SUB, line=1.4))])

    # ================= 17 보고서 평가 =================
    p += 1
    sl = content_slide(prs, CH3, "구조를 바꾸자 53건 모두 생성, 토큰은 54% 감소",
                       "합성 53건 · LLM Judge 절대+쌍대", p,
                       source="출처: docs/eval/01_보고서작성에이전트 (Judge gpt-5.6-luna, 2026-09-08)")
    strip(sl, BODY_Y + 0.05, [("47 → 53", "유효 답안", "53건 중"),
                              ("18 → 0", "외부 재시도", "생성 실패 6 → 0"),
                              ("−54.4%", "생성 토큰", "1,001만 → 457만"),
                              ("29 / 47", "쌍대 우세", "개선 전 우세 1건")], featured=2)
    y2 = BODY_Y + 1.45
    cw_ = 8.0
    eyebrow(sl, M, y2, "보고서 유형별 LLM Judge 점수 (100점 만점 · 공통 43건)")
    add_chart(sl, COL, M - 0.15, y2 + 0.35, cw_, 2.55, ["미팅", "일일", "주간", "월간"],
              [("개선 전", [97.74, 86.88, 72.50, 66.25]), ("개선 후", [96.57, 95.16, 87.08, 76.25])],
              label_fmt="0.0", gap=70, y_min=60, y_max=105, colors=(BLUE_PALE, BLUE))
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    eyebrow(sl, rx, y2, "방법과 한계", w=rw)
    for i, t in enumerate(["루브릭 5개 · 익명 순서 교차",
                           "인용 기계 검증",
                           "기간 보고서 = 단계별 참조",
                           "한계 — 시나리오 1개 · 월간 1건",
                           "한계 — 생성·Judge 동일 모델",
                           "한계 — 미팅 점수·사실 정확성 하락"]):
        write(sl, rx, y2 + 0.45 + i * 0.42, rw, 0.4,
              [("· " + t, dict(size=11, weight=400, color=SUB, line=1.35))])

    # ================= 18 영업·계약관리 (레이아웃) =================
    p += 1
    sl = content_slide(prs, CH3, "[영업·계약관리 평가 결론 입력 예정]", "[평가 방법 한 줄 입력 예정]", p)
    w4, xs4 = grid(4, CW, 0.28)
    for i in range(4):
        placeholder(sl, xs4[i], BODY_Y + 0.1, w4, 1.1, "KPI")
    placeholder(sl, M, BODY_Y + 1.45, 8.0, 2.9, "차트 또는 테스트 계층 도식")
    placeholder(sl, M + 8.35, BODY_Y + 1.45, CW - 8.35, 2.9, "평가 데이터 · 방법 · 한계")

    # ================= 19 일정관리 (레이아웃) =================
    p += 1
    sl = content_slide(prs, CH3, "[일정관리 평가 결론 입력 예정]", "[평가 방법 한 줄 입력 예정]", p)
    w3, xs3 = grid(3, CW, 0.3)
    for i in range(3):
        placeholder(sl, xs3[i], BODY_Y + 0.1, w3, 1.1, "KPI")
    w2, xs2 = grid(2, CW, 0.4)
    placeholder(sl, xs2[0], BODY_Y + 1.45, w2, 2.9, "Before — 개선 전 동작")
    placeholder(sl, xs2[1], BODY_Y + 1.45, w2, 2.9, "After — 개선 후 동작")

    # ================= 20 문서요약 =================
    p += 1
    sl = content_slide(prs, CH3, "문서요약 자동 평가 0.846, 사람 검수는 남았다",
                       "RAGEval 골든셋 10건 · 6개 지표", p,
                       source="출처: docs/document-summary-evaluation.md, output/evals/document-summary-rageval (자동 10건)")
    cw_ = 8.6
    eyebrow(sl, M, BODY_Y + 0.05, "지표별 점수 (0–1 · 높을수록 좋음)")
    add_chart(sl, BAR, M - 0.1, BODY_Y + 0.4, cw_, 3.85,
              ["overall", "completeness", "groundedness", "hallucination", "retrieval_relevance", "irrelevance"],
              [("점수", [0.846, 0.828, 0.914, 0.942, 0.892, 0.901])], label_fmt="0.000", gap=60,
              y_min=0.0, y_max=1.0)
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    eyebrow(sl, rx, BODY_Y + 0.05, "평가 흐름", w=rw)
    for i, t in enumerate(["합성 문서 7개", "QRA 자동 골든셋 10건", "요약 6 · 검색 4 Judge 채점"]):
        circle(sl, rx, BODY_Y + 0.5 + i * 0.62, 0.28, str(i + 1), fill=TINT, color=BLUE_600, size=10)
        write(sl, rx + 0.4, BODY_Y + 0.52 + i * 0.62, rw - 0.4, 0.45,
              [(t, dict(size=11.5, weight=500, color=INK, line=1.35))])
    line(sl, rx, BODY_Y + 2.5, rx + rw, BODY_Y + 2.5, color=BORDER)
    eyebrow(sl, rx, BODY_Y + 2.7, "읽을 때 주의", w=rw, color=MUTED)
    for i, t in enumerate(["요약 6건 0.91–0.98", "사람 검수 pending", "합성 문서 기준"]):
        write(sl, rx, BODY_Y + 3.15 + i * 0.4, rw, 0.38,
              [("· " + t, dict(size=11, weight=400, color=SUB, line=1.35))])

    # ================= 21 OCR =================
    p += 1
    sl = content_slide(prs, CH3, "명함·PDF는 필드 100%, 사업자등록증은 80%",
                       "유형별 정답표 대조 · 오류 필드는 사용자 검토", p,
                       source="출처: 자료요약 OCR 평가표, RunPod OCR 워커 계약 문서 (유형별 1건 소표본)")
    strip(sl, BODY_Y + 0.05, [("7 / 7", "명함 필드", "구조화 6 / 6"),
                              ("100%", "취업규칙 PDF 35쪽", "장 제목 12/12"),
                              ("12 / 15", "사업자등록증 필드", "80% · 3개 불일치")])
    y2 = BODY_Y + 1.45
    cw_ = 6.6
    eyebrow(sl, M, y2, "RunPod 웜 스타트 처리 시간 (초 · 각 5회)")
    add_chart(sl, COL, M - 0.15, y2 + 0.35, cw_, 2.55, ["명함 이미지", "PDF 문서"],
              [("평균 처리 시간(초)", [2.24, 1.23])], label_fmt="0.00", gap=140, y_max=3.0)
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    eyebrow(sl, rx, y2, "틀린 필드와 대응", w=rw)
    for i, t in enumerate(["법인등록번호 · 법인명 · 개업연월일 불일치",
                           "자동 승인 제외 → 사용자 검토",
                           "유형별 1건 — 전체 정확도 아님",
                           "손글씨 오류율 45% — 별도 과제"]):
        write(sl, rx, y2 + 0.45 + i * 0.62, rw, 0.6,
              [("· " + t, dict(size=11, weight=400, color=SUB, line=1.4))])

    # ================= 22 RAG =================
    p += 1
    sl = content_slide(prs, CH3, "RAG 답변 4건 중 1건은 옛 계약서 날짜를 답했다",
                       "요약보다 검색 답변의 편차가 큼 — 원인은 최신성", p,
                       source="출처: output/evals/document-summary-rageval/llm_judge_results.json (요약 6 · 검색 4)")
    cw_ = 8.3
    eyebrow(sl, M, BODY_Y + 0.05, "케이스별 overall 점수 (0–1)")
    add_chart(sl, COL, M - 0.15, BODY_Y + 0.4, cw_, 3.7,
              ["요약1", "요약2", "요약3", "요약4", "요약5", "요약6", "검색1", "검색2", "검색3", "검색4"],
              [("overall", [0.97, 0.91, 0.96, 0.98, 0.93, 0.96, 0.25, 0.78, 0.90, 0.82])],
              label_fmt="0.00", gap=50, y_max=1.0, label_size=9)
    rx = M + cw_ + 0.45
    rw = CW - cw_ - 0.45
    card(sl, rx, BODY_Y + 0.05, rw, 2.35, fill=TINT, shadow=False)
    write(sl, rx + 0.25, BODY_Y + 0.3, rw - 0.5, 0.3,
          [("실패 사례 — 검색1 (0.25)", dict(size=13, weight=700, color=BLUE_DEEP, line=1.25))])
    for i, (k, v) in enumerate([("질문", "현재 기준 설치 예정일"), ("기대", "2026-10-08 추가 자료"),
                                ("답변", "2026-09-24 기존 계약서")]):
        y = BODY_Y + 0.75 + i * 0.44
        write(sl, rx + 0.25, y, 0.6, 0.3, [(k, dict(size=10.5, weight=600, color=MUTED, line=1.3))])
        write(sl, rx + 0.9, y, rw - 1.15, 0.3, [(v, dict(size=11, weight=500, color=INK, line=1.3))])
    write(sl, rx + 0.25, BODY_Y + 2.05, rw - 0.5, 0.3,
          [("→ 최신 자료 대신 옛 계약서 참조", dict(size=11, weight=600, color=BLUE_600, line=1.3))])
    eyebrow(sl, rx, BODY_Y + 2.65, "읽을 때 주의", w=rw, color=MUTED)
    for i, t in enumerate(["평가 스크립트 = 단어 겹침 검색 (운영은 키워드+벡터)",
                           "보고서 RAG · 브리핑은 별도 평가 없음"]):
        write(sl, rx, BODY_Y + 3.1 + i * 0.62, rw, 0.6,
              [("· " + t, dict(size=11, weight=400, color=SUB, line=1.4))])

    # ================= 23 실패와 개선 (레이아웃) =================
    p += 1
    sl = content_slide(prs, CH3, "[실패와 개선 결론 입력 예정]", "[부제 입력 예정]", p)
    w4, xs4 = grid(4, CW, 0.25)
    for r_i in range(2):
        y = BODY_Y + 0.15 + r_i * 2.15
        write(sl, M, y - 0.05, 2.0, 0.28, [("사례 " + str(r_i + 1), dict(size=11, weight=700, color=BLUE_600, line=1.2))])
        for c_i, lb in enumerate(["처음 접근", "무엇이 문제였나", "어떻게 바꿨나", "결과"]):
            placeholder(sl, xs4[c_i], y + 0.28, w4, 1.6, lb)

    # ================= 24 보완점 (레이아웃) =================
    p += 1
    sl = content_slide(prs, CH3, "[보완점 결론 입력 예정]", "[부제 입력 예정]", p)
    w3, xs3 = grid(3, CW, 0.3)
    for i in range(3):
        placeholder(sl, xs3[i], BODY_Y + 0.15, w3, 3.0, "한계 " + str(i + 1) + " — 영역 · 내용 · 영향")
    placeholder(sl, M, BODY_Y + 3.35, CW, 0.85, "So What — 한 줄 정리")

    # ================= 25 향후 계획 (레이아웃) =================
    p += 1
    sl = content_slide(prs, CH3, "[향후 계획 결론 입력 예정]", "[부제 입력 예정]", p)
    line(sl, M + 0.6, BODY_Y + 0.9, M + CW - 0.6, BODY_Y + 0.9, color=BORDER, width=1.5)
    w3, xs3 = grid(3, CW, 0.5)
    for i, lb in enumerate(["단기", "중기", "장기"]):
        x = xs3[i]
        circle(sl, x + w3 / 2 - 0.22, BODY_Y + 0.68, 0.44, str(i + 1), fill=BLUE if i == 0 else TINT,
               color=WHITE if i == 0 else BLUE_600, size=12.5)
        write(sl, x, BODY_Y + 1.3, w3, 0.3, [(lb + " 계획", dict(size=15, weight=700, color=INK, line=1.25))], align=C)
        placeholder(sl, x, BODY_Y + 1.75, w3, 2.4, PH)

    # ================= 26 결론 =================
    p += 1
    sl = content_slide(prs, CH3, "담당자의 머릿속을 조직의 기록으로 남긴다",
                       "인수인계 질문 11개 영역 = 데이터", p,
                       source="출처: SalesLuv 데이터 구조(develop, 2026-09-17)")
    areas = [("고객 관계 맥락", "회사 · 담당자 · 부서"),
             ("접점 히스토리", "일정 · 활동 · 보고서"),
             ("고객의 실제 발언", "원문 · STT · 첨부"),
             ("발언의 대상·범위", "공통 · 딜별 + 근거 ID"),
             ("구매 판단 신호", "권한 · 경쟁사 등 13개 특성"),
             ("딜 진행 맥락", "단계 · 금액 · 계약 · 납기"),
             ("과거 맥락", "이전 확정 보고서"),
             ("이슈·불만 맥락", "불만 · 긴급도 · 처리"),
             ("위험과 다음 행동", "만료 · 지연 · 추천"),
             ("문서 지식", "원문 추출 · 요약 · RAG"),
             ("판단의 책임·이력", "작성자 · revision · 승인")]
    w4, xs4 = grid(4, CW, 0.22)
    th_ = 1.0
    gap_ = 0.2
    for i, (t, d) in enumerate(areas):
        x = xs4[i % 4]
        y = BODY_Y + 0.15 + (i // 4) * (th_ + gap_)
        card(sl, x, y, w4, th_, fill=WHITE, line=BORDER)
        write(sl, x + 0.22, y + 0.2, w4 - 0.44, 0.3, [(t, dict(size=13, weight=700, color=INK, line=1.25))])
        write(sl, x + 0.22, y + 0.56, w4 - 0.44, 0.3, [(d, dict(size=10.5, weight=400, color=SUB, line=1.3))])
    gx, gy = xs4[3], BODY_Y + 0.15 + 2 * (th_ + gap_)
    gradient_card(sl, gx, gy, w4, th_, radius=0.16)
    write(sl, gx + 0.25, gy + 0.2, w4 - 0.5, 0.65,
          [("AI는 구조화 · 요약 · 제안까지", dict(size=12.5, weight=700, color=WHITE, line=1.35)),
           ("확정은 사람의 승인으로", dict(size=12.5, weight=700, color=BLUE_PALE, line=1.35, spacing_before=2))])

    # ================= 27 Q&A =================
    p += 1
    sl = blank(prs, dark=True)
    frame(sl, page=p, dark=True)
    write(sl, M, 3.0, CW, 1.2, [("Q & A", dict(size=54, weight=700, color=WHITE, line=1.15))], align=C)
    write(sl, M, 4.3, CW, 0.5,
          [("감사합니다 · Thank You", dict(size=17, weight=500, color=RGBColor(0xC9, 0xD3, 0xE0), line=1.4))],
          align=C)

    # ================= 부록 1 ERD =================
    p += 1
    sl = content_slide(prs, "부록", "모든 데이터는 딜을 중심으로 연결된다", "핵심 20개 엔티티 요약", p,
                       source="출처: backend/app/models/ (develop, 2026-09-17)")
    cx, cy, cwd, chd = 5.0, BODY_Y + 1.55, 3.33, 0.85
    hub = card(sl, cx, cy, cwd, chd, fill=BLUE, shadow=True)
    tf = hub.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, "sales_deal", size=17, weight=700, color=WHITE, line=1.2, align=C, first=True)
    para(tf, "영업 딜", size=11, weight=400, color=BLUE_PALE, line=1.3, align=C)
    groups = [("조직 · 권한", "team · member", M, BODY_Y + 0.1),
              ("고객", "customer_company · customer_contact", M + 4.35, BODY_Y + 0.1),
              ("영업 실행", "sales_pipeline(+stage) · product\nsales_deal_item · purchase_order(+item)", M + 8.7, BODY_Y + 0.1),
              ("활동 · 보고", "activity · report\nreport_deal · report_submission", M, BODY_Y + 3.05),
              ("AI 실행", "agent_run\ncontract_next_meeting_suggestion", M + 4.35, BODY_Y + 3.05),
              ("문서 · C/S", "document · file\ndocument_chunk · support_request", M + 8.7, BODY_Y + 3.05)]
    gw, gh = 4.1, 1.15
    bus = cy + chd / 2
    line(sl, M + gw / 2, BODY_Y + gh, M + gw / 2, BODY_Y + 3.05, color=BLUE_PALE, width=1.2)
    line(sl, M + 8.7 + gw / 2, BODY_Y + gh, M + 8.7 + gw / 2, BODY_Y + 3.05, color=BLUE_PALE, width=1.2)
    line(sl, M + gw / 2, bus, cx, bus, color=BLUE_PALE, width=1.2)
    line(sl, cx + cwd, bus, M + 8.7 + gw / 2, bus, color=BLUE_PALE, width=1.2)
    line(sl, M + 4.35 + gw / 2, BODY_Y + gh, M + 4.35 + gw / 2, cy, color=BLUE_PALE, width=1.2)
    line(sl, M + 4.35 + gw / 2, cy + chd, M + 4.35 + gw / 2, BODY_Y + 3.05, color=BLUE_PALE, width=1.2)
    for label, items, x, y in groups:
        card(sl, x, y, gw, gh, fill=WHITE, line=BORDER)
        write(sl, x + 0.22, y + 0.18, gw - 0.44, 0.28, [(label, dict(size=12.5, weight=700, color=BLUE_600, line=1.25))])
        lines = [(t, dict(size=10.5, weight=400, color=SUB, line=1.4, spacing_before=(0 if i == 0 else 2)))
                 for i, t in enumerate(items.split("\n"))]
        write(sl, x + 0.22, y + 0.52, gw - 0.44, 0.55, lines)

    # ================= 부록 2 WBS =================
    p += 1
    sl = content_slide(prs, "부록", "7주 동안 기획에서 배포 검증까지 진행했다", "주차별 주요 작업", p,
                       source="출처: 프로젝트 WBS")
    weeks = [("1주", "요구사항 분석 · 기획"), ("2주", "DB 기반 구축 · 화면 설계"),
             ("3주", "FE/BE 개발 · 중간 발표"), ("4주", "데이터 전처리 · AI 모델링"),
             ("5주", "모델 평가 · 기능 통합"), ("6주", "통합 테스트 · 배포 검증"),
             ("7주", "산출물 검수 · 최종 발표")]
    n = len(weeks)
    w_ = (CW - 0.2 * (n - 1)) / n
    ry = BODY_Y + 2.05
    line(sl, M + w_ / 2, ry, M + CW - w_ / 2, ry, color=BORDER, width=1.5)
    for i, (wk, t) in enumerate(weeks):
        x = M + i * (w_ + 0.2)
        circle(sl, x + w_ / 2 - 0.09, ry - 0.09, 0.18, "", fill=(BLUE if i >= n - 2 else BLUE_PALE))
        write(sl, x, ry - 0.75, w_, 0.35, [(wk, dict(size=15, weight=700, color=INK, line=1.2))], align=C)
        write(sl, x, ry + 0.35, w_, 1.2, [(t, dict(size=11, weight=400, color=SUB, line=1.45))], align=C)

    prs.save(OUT)
    print("saved", OUT, len(prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    build()
