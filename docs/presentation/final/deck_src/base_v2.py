"""SalesLuv 최종 발표 v2 — 에디토리얼 대형 타이포 디자인 시스템.

v1(base.py)과 다른 점
  - 본문 최소 18pt / 제목 34pt / 주석 14pt 로 타입 스케일을 키움
  - 카드는 테두리·그림자 없는 플랫 틴트 블록, 강조는 왼쪽 액센트 바
  - 번호는 작은 원형 뱃지 대신 큰 고스트 숫자
  - 표지·파트 디바이더·Q&A 는 다크 배경으로 통일
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from base import (  # 재사용: 순수 프리미티브
    BAR, BLUE, BLUE_500, BLUE_600, BLUE_DEEP, BLUE_LIGHT, BLUE_PALE, BORDER, COL, DARK,
    INK, MUTED, SUB, WHITE, add_chart, arrow, blank, card, circle, line, new_deck, para,
    pill, textbox, write,
)

# ---------- 추가 색 ----------
TINT = RGBColor(0xF3, 0xF6, 0xFE)      # 연한 브랜드 틴트
TINT2 = RGBColor(0xE6, 0xEE, 0xFD)     # 한 단계 진한 틴트
GRAY = RGBColor(0xF5, 0xF6, 0xF8)
DARK_SUB = RGBColor(0xC4, 0xCF, 0xDD)
HAIRLINE = RGBColor(0xE6, 0xE8, 0xEC)

L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT

# ---------- 지오메트리 ----------
SW, SH = 13.333, 7.5
M = 0.62                       # 안전 여백 (요구 0.45 이상)
CW = SW - 2 * M                # 12.093
EYE_Y = 0.45
H1_Y = 0.96
SUB_Y = 1.70
RULE_Y = 2.34
TOP = 2.56                     # 부제 있을 때 본문 시작
TOP_NOSUB = 2.02               # 부제 없을 때
BOT = 6.62                     # 본문 하단
FOOT_Y = 6.86

LOGO_W, LOGO_H = 1.15, 0.268
HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = "/Volumes/Jexists/skn30/SKN30-FINAL-1Team-dev/frontend/src/assets/full-logo.png"
LOGO_WHITE = os.path.join(HERE, "assets", "full-logo-white.png")

# ---------- 타입 스케일 ----------
T_H1 = 34
T_SUB = 19
T_EYE = 14
T_CARD = 21
T_BODY = 18
T_NOTE = 14
T_KPI = 42


def make_white_logo():
    from PIL import Image
    os.makedirs(os.path.dirname(LOGO_WHITE), exist_ok=True)
    img = Image.open(LOGO).convert("RGBA")
    img.putdata([(255, 255, 255, a) for (_, _, _, a) in img.getdata()])
    img.save(LOGO_WHITE)


# ---------- 프레임 ----------
def _bar(sl, x, y, w, h, color):
    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def slide(prs, page, *, chapter=None, headline=None, subtitle=None, source=None, dark=False):
    """표준 본문 슬라이드. (slide, 본문 시작 y) 반환."""
    sl = blank(prs, dark=dark)
    fg_eye = WHITE if dark else BLUE
    if chapter:
        _bar(sl, M, EYE_Y + 0.035, 0.085, 0.2, fg_eye)
        write(sl, M + 0.22, EYE_Y, 7.0, 0.3,
              [(chapter, dict(size=T_EYE, weight=600, color=fg_eye, line=1.25))])
    sl.shapes.add_picture(LOGO_WHITE if dark else LOGO,
                          Inches(SW - M - LOGO_W), Inches(EYE_Y - 0.02),
                          Inches(LOGO_W), Inches(LOGO_H))
    if headline:
        write(sl, M, H1_Y, CW, 0.7,
              [(headline, dict(size=T_H1, weight=700, color=(WHITE if dark else INK), line=1.2))])
    top = TOP
    if subtitle:
        write(sl, M, SUB_Y, CW, 0.4,
              [(subtitle, dict(size=T_SUB, weight=500, color=(DARK_SUB if dark else SUB), line=1.4))])
        line(sl, M, RULE_Y, M + CW, RULE_Y, color=HAIRLINE)
    else:
        line(sl, M, RULE_Y - 0.52, M + CW, RULE_Y - 0.52, color=HAIRLINE)
        top = TOP_NOSUB
    if source:
        write(sl, M, FOOT_Y, CW - 1.2, 0.3,
              [(source, dict(size=T_NOTE, weight=400, color=MUTED, line=1.35))])
    write(sl, SW - M - 1.0, FOOT_Y - 0.03, 1.0, 0.3,
          [("%02d" % page, dict(size=16, weight=600, color=(BLUE_PALE if not dark else SUB), line=1.25))],
          align=R)
    return sl, top


def divider(prs, number, title, lead, page):
    sl = blank(prs, dark=True)
    sl.shapes.add_picture(LOGO_WHITE, Inches(SW - M - LOGO_W), Inches(EYE_Y - 0.02),
                          Inches(LOGO_W), Inches(LOGO_H))
    write(sl, M, 1.55, 6.0, 1.6, [(number, dict(size=120, weight=700, color=BLUE, line=1.0))])
    _bar(sl, M, 3.62, 1.5, 0.075, WHITE)
    write(sl, M, 3.95, 10.0, 1.0, [(title, dict(size=48, weight=700, color=WHITE, line=1.15))])
    write(sl, M, 5.05, 10.0, 0.5, [(lead, dict(size=20, weight=500, color=DARK_SUB, line=1.4))])
    write(sl, SW - M - 1.0, FOOT_Y - 0.03, 1.0, 0.3,
          [("%02d" % page, dict(size=16, weight=600, color=SUB, line=1.25))], align=R)
    return sl


# ---------- 콘텐츠 프리미티브 ----------
def label(sl, x, y, text, *, w=6.0, color=BLUE):
    """섹션 라벨 (15pt)."""
    _bar(sl, x, y + 0.045, 0.075, 0.18, color)
    write(sl, x + 0.2, y, w, 0.28, [(text, dict(size=15, weight=600, color=color, line=1.25))])


def flat(sl, x, y, w, h, *, fill=TINT, accent=None, radius=0.12):
    """테두리·그림자 없는 플랫 카드. accent 색을 주면 왼쪽 액센트 바."""
    sh = card(sl, x, y, w, h, fill=fill, radius=radius, shadow=False)
    if accent is not None:
        _bar(sl, x, y, 0.075, h, accent)
    return sh


def ghost(sl, x, y, w, text, *, size=44, color=BLUE_PALE, align=L):
    write(sl, x, y, w, size / 72 * 1.15 + 0.1,
          [(text, dict(size=size, weight=700, color=color, line=1.1))], align=align)


def numcard(sl, x, y, w, h, num, title, desc, *, active=False, num_size=40, title_size=T_CARD,
            body_size=T_BODY, title_lines=1, pad=0.3):
    fill = TINT2 if active else TINT
    flat(sl, x, y, w, h, fill=fill, accent=(BLUE if active else None))
    px = x + (pad + 0.08 if active else pad)
    iw = w - (px - x) - pad
    ghost(sl, px, y + 0.22, iw, num, size=num_size, color=(BLUE if active else BLUE_LIGHT))
    ty = y + 0.22 + num_size / 72 * 1.15 + 0.12
    write(sl, px, ty, iw, 0.35,
          [(title, dict(size=title_size, weight=600, color=INK, line=1.25))])
    if desc:
        write(sl, px, ty + title_size / 72 * 1.25 * title_lines + 0.12, iw, h,
              [(desc, dict(size=body_size, weight=400, color=SUB, line=1.45))])


def kpi(sl, x, y, w, value, lab, sub=None, *, color=BLUE, size=T_KPI, align=L, tint=False, h=1.55):
    if size == T_KPI and len(value) > 6:   # "47 → 53" 처럼 긴 값은 줄바꿈 대신 축소
        size = 34
    if tint:
        flat(sl, x, y, w, h)
        x += 0.32
        w -= 0.64
        y += 0.2
    write(sl, x, y, w, size / 72 * 1.2 + 0.1,
          [(value, dict(size=size, weight=700, color=color, line=1.1))], align=align)
    yy = y + size / 72 * 1.2 + 0.08
    write(sl, x, yy, w, 0.34, [(lab, dict(size=17, weight=600, color=INK, line=1.3))], align=align)
    if sub:
        write(sl, x, yy + 0.36, w, 0.3,
              [(sub, dict(size=T_NOTE, weight=400, color=MUTED, line=1.3))], align=align)


def bullets(sl, x, y, w, items, *, size=T_BODY, gap=0.52, color=SUB, marker=BLUE_LIGHT, weight=400):
    for i, t in enumerate(items):
        yy = y + i * gap
        _bar(sl, x, yy + 0.09, 0.09, 0.09, marker)
        write(sl, x + 0.26, yy, w - 0.26, gap,
              [(t, dict(size=size, weight=weight, color=color, line=1.4))])
    return y + len(items) * gap


def keyline(sl, y, text, *, x=M, w=CW, size=T_BODY):
    """결론 한 줄 강조 띠."""
    h = 0.62
    flat(sl, x, y, w, h, fill=TINT2, accent=BLUE)
    write(sl, x + 0.36, y + (h - size / 72 * 1.35) / 2, w - 0.6, h,
          [(text, dict(size=size, weight=600, color=BLUE_DEEP, line=1.35))])


def placeholder(sl, x, y, w, h, text, *, size=T_BODY):
    from pptx.oxml import parse_xml
    sh = card(sl, x, y, w, h, fill=GRAY, shadow=False, line=BORDER)
    sh.line._get_or_add_ln().append(parse_xml(
        '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, weight=500, color=MUTED, line=1.4, align=C, first=True)
    return sh


def note(sl, x, y, w, text, *, size=T_NOTE, color=MUTED, align=L, weight=400):
    return write(sl, x, y, w, 0.6, [(text, dict(size=size, weight=weight, color=color, line=1.4))],
                 align=align)


def grid(n, gap, *, total=CW, x0=M):
    w = (total - gap * (n - 1)) / n
    return w, [x0 + i * (w + gap) for i in range(n)]


def table(sl, x, y, w, h, rows, *, col_w=None, highlight_col=None, highlight_row=None,
          size=T_BODY, head_size=16):
    """최대 5행 × 4열 규칙에 맞춘 표."""
    nr, nc = len(rows), len(rows[0])
    assert nr <= 5 and nc <= 4, "표는 최대 5행 x 4열"
    tbl = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_w:
        for i, cw_ in enumerate(col_w):
            tbl.columns[i].width = Inches(cw_)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(h / nr)
        for ci, val in enumerate(row):
            cell = tbl.cell(r, ci)
            cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.16)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            hl = (highlight_col is not None and ci == highlight_col) or \
                 (highlight_row is not None and r == highlight_row)
            para(cell.text_frame, val,
                 size=head_size if r == 0 else size,
                 weight=600 if (r == 0 or ci == 0 or hl) else 400,
                 color=(WHITE if r == 0 else (BLUE if hl else (INK if ci == 0 else SUB))),
                 line=1.25, first=True, align=L if ci == 0 else C)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = BLUE_DEEP
            elif hl:
                cell.fill.fore_color.rgb = TINT2
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else GRAY
    return tbl
