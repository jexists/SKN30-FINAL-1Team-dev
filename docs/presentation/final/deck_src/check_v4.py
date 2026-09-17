"""v4 검수 — 수치 출처 · 페이지 번호 · 레이아웃을 자동으로 확인한다.

새로 넣은 30 · 31 · 32p 에 적힌 숫자가 각자의 근거 평가 pptx 에 실제로 있는지 대조한다.
슬라이드마다 근거 파일을 하나로 못박아, 다른 평가의 숫자가 섞이면 실패한다.
"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches

from build_v4 import AT, DATA_X, DATA_W, KPI_H, KPI_Y, OUT, ROW_H, ROW_Y
from textwidth import width
from base_v3 import BOT, CW, M, SW, TOP

EVAL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "eval")
SOURCES = {
    AT + 1: "SalesLuv_미팅브리핑_평가.pptx",
    AT + 2: "SalesLuv_보고서작성에이전트_평가_4p.pptx",
    AT + 3: "SalesLuv_딜승산예측모델_평가.pptx",
}
NUM = re.compile(r"\d[\d,.]*")
# 근거 덱에 숫자로 적혀 있지 않고 본문 서술에서 끌어온 값. 확인한 근거를 함께 적는다.
DERIVED = {
    AT + 1: {"1": "1 · 3 · 5 · 10 · 30회 — 1p 평가 축",
             "3": "1 · 3 · 5 · 10 · 30회 — 1p 평가 축",
             "5": "1~5회 시점 평균 · 1p 평가 축"},
    AT + 2: {"12": "1p '일일 12 · 주간 4 · 월간 1'",
             "4": "1p '일일 12 · 주간 4 · 월간 1'",
             "1": "1p '일일 12 · 주간 4 · 월간 1'"},
    AT + 3: {"4": "2p '행마다 Unknown 을 4개로 맞춘다'",
             "10": "2p '448건 × 10세트'",
             "7": "5p 후보 7종 (RF 튜닝 … Stacking_LR)",
             "2": "5p 'Brier 1위 · FP 2위'",
             "1": "5p 'Brier 1위 · FP 2위'"},
}


def texts(path):
    """평가 덱의 모든 텍스트 + 차트 값을 한 덩어리로 모은다."""
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif getattr(sh, "has_chart", False):
                ch = sh.chart
                out.extend(str(c) for c in ch.plots[0].categories)
                for ser in ch.plots[0].series:
                    for pt in ser._element.iter():
                        if pt.tag.endswith("}v") and pt.text:
                            out.append(pt.text)
            elif sh.has_table:
                out.extend(c.text for r in sh.table.rows for c in r.cells)
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)
    for s in Presentation(path).slides:
        walk(s.shapes)
    return "\n".join(out)


def canon(tok):
    """'10.01' 과 '10.01M', '4,480' 과 '4480' 이 같게 보이도록 정규화."""
    t = tok.replace(",", "")
    if "." in t:
        t = t.rstrip("0").rstrip(".")
    return t


def numbers(blob):
    got = set()
    for tok in NUM.findall(blob):
        got.add(canon(tok))
        if "." in tok:                       # 0.18797 ↔ 0.1880
            for nd in (3, 4):
                got.add(canon("%.*f" % (nd, float(tok.replace(",", "")))))
    return got


def slide_text(sl):
    out = []
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.append(sh.text_frame.text)
    return "\n".join(out)


def boxes(sl):
    return [(sh, sh.left / 914400, sh.top / 914400,
             (sh.width or 0) / 914400, (sh.height or 0) / 914400) for sh in sl.shapes]


def main():
    prs = Presentation(OUT)
    slides = list(prs.slides)
    fail = []

    # ---------- 1. 슬라이드 수 ----------
    assert len(slides) == 59, "슬라이드 %d 장 (59 이어야 함)" % len(slides)

    # ---------- 2. 페이지 번호 ----------
    for i, sl in enumerate(slides, 1):
        for sh in sl.shapes:
            if not sh.has_text_frame or sh.top is None:
                continue
            if abs(sh.left - Inches(SW - M - 1.0)) < Inches(0.05) and \
               abs(sh.top - Inches(6.86 - 0.03)) < Inches(0.05):
                got = sh.text_frame.text.strip()
                if got != "%02d" % i:
                    fail.append("페이지 번호 %d 장 = %r" % (i, got))

    # ---------- 3. 수치 대조 ----------
    for page, fname in SOURCES.items():
        blob = texts(os.path.join(EVAL, fname))
        allowed = numbers(blob)
        sl = slides[page - 1]
        body = slide_text(sl)
        body = body.replace("출처: SalesLuv 보고서작성 에이전트 평가", "")   # 파일명 속 '4p'
        for tok in NUM.findall(body):
            c = canon(tok)
            if c in allowed:
                continue
            if c in DERIVED.get(page, {}):
                continue
            if c == str(page):                # 페이지 번호
                continue
            if re.fullmatch(r"0\d", tok):      # 01~09 = 챕터 · 단계 장식 번호
                continue
            fail.append("%dp 의 %r 을 %s 에서 찾지 못함" % (page, tok, fname))

    # ---------- 4. 다른 평가의 숫자가 섞였는지 ----------
    keys = {AT + 1: ["89.8", "96.3", "180", "1.7만"],
            AT + 2: ["95.4", "10.01", "4.57", "76.3"],
            AT + 3: ["0.1880", "16.15", "4.01", "4,480"]}
    for page in keys:
        for other, vals in keys.items():
            if other == page:
                continue
            for v in vals:
                if v in slide_text(slides[page - 1]):
                    fail.append("%dp 에 %dp 의 값 %r 이 있다" % (page, other, v))

    # ---------- 5. 레이아웃 ----------
    for page in SOURCES:
        sl = slides[page - 1]
        for sh, x, y, w, h in boxes(sl):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:      # 로고
                continue
            if y > BOT:                                      # 출처 · 페이지 번호
                continue
            if x < M - 0.01 or x + w > M + CW + 0.01:
                fail.append("%dp 좌우 여백 벗어남: %.2f~%.2f" % (page, x, x + w))
            if y + h > BOT + 0.01 and sh.has_text_frame and sh.text_frame.text.strip():
                fail.append("%dp 본문 하단 넘침: %.2f" % (page, y + h))
        # 한 줄로 끝나야 하는 글상자가 칸을 넘지 않는가 (넘으면 줄바꿈 → 카드 밖으로 샌다)
        for sh, x, y, w, h in boxes(sl):
            if not (sh.has_text_frame and sh.text_frame.text.strip()) or not sh.width:
                continue
            if h > 0.36 or y > BOT:            # 여러 줄을 허용한 칸 · 푸터는 제외
                continue
            runs = sh.text_frame.paragraphs[0].runs
            if not runs:
                continue
            pt = runs[0].font.size.pt if runs[0].font.size else 15
            wt = 700 if runs[0].font.bold else 400
            txt = sh.text_frame.text.strip()
            if width(txt, pt, wt) > w + 0.01:
                fail.append("%dp 한 줄 넘침: %r (%.2f > %.2f in)"
                            % (page, txt[:26], width(txt, pt, wt), w))

        # 글자끼리 겹치지 않는가 (작은 쪽 면적의 30% 넘게 물리면 실패)
        tb = [(sh.text_frame.text.strip()[:18], x, y, w, h) for sh, x, y, w, h in boxes(sl)
              if sh.has_text_frame and sh.text_frame.text.strip() and sh.width]
        for i, (t1, x1, y1, w1, h1) in enumerate(tb):
            for t2, x2, y2, w2, h2 in tb[i + 1:]:
                ox = min(x1 + w1, x2 + w2) - max(x1, x2)
                oy = min(y1 + h1, y2 + h2) - max(y1, y2)
                if ox <= 0 or oy <= 0:
                    continue
                if ox * oy > 0.3 * min(w1 * h1, w2 * h2):
                    fail.append("%dp 글자 겹침: %r ↔ %r" % (page, t1, t2))

        # 블록끼리 겹치지 않는가
        bands = [("KPI", KPI_Y, KPI_Y + KPI_H), ("본문", ROW_Y, ROW_Y + ROW_H)]
        for name, a, b in bands:
            if a < TOP - 0.01 or b > BOT + 0.01:
                fail.append("%dp %s 밴드가 본문 영역 밖: %.2f~%.2f" % (page, name, a, b))
        if KPI_Y + KPI_H > ROW_Y:
            fail.append("%dp KPI 띠와 본문이 겹친다" % page)
        if M + 7.20 > DATA_X or DATA_X + DATA_W > M + CW + 0.01:
            fail.append("%dp 좌우 2단이 겹치거나 넘친다" % page)

    if fail:
        print("실패 %d 건" % len(fail))
        for f in fail:
            print("  -", f)
        sys.exit(1)
    print("통과 — 59장 · 페이지 번호 · 수치 출처 · 레이아웃 이상 없음")
    for page, fname in SOURCES.items():
        print("  %02dp ← %s" % (page, fname))


if __name__ == "__main__":
    main()
