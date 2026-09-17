"""Brandlogy 디자인 시스템 공통 프레임 (16:9, Pretendard)."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------- 색 ----------
BLUE = RGBColor(0x14, 0x56, 0xF0)
BLUE_500 = RGBColor(0x3B, 0x82, 0xF6)
BLUE_LIGHT = RGBColor(0x60, 0xA5, 0xFA)
BLUE_PALE = RGBColor(0xBF, 0xDB, 0xFE)
BLUE_DEEP = RGBColor(0x17, 0x43, 0x7D)
BLUE_600 = RGBColor(0x25, 0x63, 0xEB)
INK = RGBColor(0x22, 0x22, 0x22)
INK_DARK = RGBColor(0x18, 0x18, 0x1B)
SUB = RGBColor(0x45, 0x51, 0x5E)
MUTED = RGBColor(0x8E, 0x8E, 0x93)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF0, 0xF0, 0xF0)
DIVIDER = RGBColor(0xF2, 0xF3, 0xF5)
BORDER = RGBColor(0xE5, 0xE7, 0xEB)
DARK = RGBColor(0x18, 0x1E, 0x25)
PINK = RGBColor(0xEA, 0x5E, 0xC1)

# ---------- 글꼴 (설치된 4개 굵기) ----------
F400 = "Pretendard"
F500 = "Pretendard Medium"
F600 = "Pretendard SemiBold"
F700 = "Pretendard"  # bold=True

# ---------- 고정 좌표 ----------
SLIDE_W, SLIDE_H = 13.333, 7.5
M = 0.5
CW = SLIDE_W - 2 * M
HEADER_Y = 0.4
HEADLINE_Y = 1.0
SUBTITLE_Y = 1.63
BODY_Y = 2.39
BODY_B = 6.85
BODY_H = BODY_B - BODY_Y
FOOT_Y = 7.05
LOGO_W, LOGO_H = 1.03, 0.24
LOGO = "/Volumes/Jexists/skn30/SKN30-FINAL-1Team-dev/frontend/src/assets/full-logo.png"
LOGO_WHITE = "/private/tmp/claude-501/-Volumes-Jexists-skn30-SKN30-FINAL-1Team-dev/e8a86c8c-cc1b-4ad0-af44-5051a3e6a99b/scratchpad/deck/logo-white.png"


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def blank(prs, dark=False):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    if dark:
        bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK
        bg.line.fill.background()
        bg.shadow.inherit = False
    return sl


# ---------- 텍스트 ----------
def textbox(sl, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, *, size=13, weight=400, color=INK, space_after=0, line=1.5,
         align=None, first=False, spacing_before=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    p.line_spacing = line
    p.space_after = Pt(space_after)
    p.space_before = Pt(spacing_before)
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.color.rgb = color
    if weight >= 700:
        f.name = F700
        f.bold = True
    elif weight == 600:
        f.name = F600
    elif weight == 500:
        f.name = F500
    else:
        f.name = F400
    return p


def write(sl, x, y, w, h, lines, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: [(text, kwargs), ...]"""
    tf = textbox(sl, x, y, w, h, align=align, anchor=anchor)
    for i, (text, kw) in enumerate(lines):
        kw = dict(kw)
        kw.setdefault("align", align)
        para(tf, text, first=(i == 0), **kw)
    return tf


# ---------- 도형 ----------
def _shadow(shape, blur=6, dist=4, alpha=0x14, color="000000"):
    spPr = shape._element.spPr
    for tag in ("a:effectLst",):
        old = spPr.find(qn(tag))
        if old is not None:
            spPr.remove(old)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur*12700)}" dist="{int(dist*12700)}" dir="5400000" '
        'rotWithShape="0">'
        f'<a:srgbClr val="{color}"><a:alpha val="{int(alpha/255*100000)}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    from pptx.oxml import parse_xml

    spPr.append(parse_xml(xml))


def card(sl, x, y, w, h, *, fill=WHITE, radius=0.13, line=None, shadow=True,
         brand_glow=False, line_w=1):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if brand_glow:
        _shadow(sh, blur=15, dist=0, alpha=0x29, color="2C1E74")
    elif shadow:
        _shadow(sh, blur=6, dist=4, alpha=0x14)
    sh.text_frame.word_wrap = True
    return sh


def gradient_card(sl, x, y, w, h, radius=0.24):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.line.fill.background()
    sh.shadow.inherit = False
    fill = sh.fill
    fill.gradient()
    fill.gradient_angle = 45.0
    stops = fill.gradient_stops
    stops[0].color.rgb = BLUE
    stops[0].position = 0.0
    stops[1].color.rgb = BLUE_LIGHT
    stops[1].position = 1.0
    # 중간 스톱 추가
    from pptx.oxml import parse_xml

    gsLst = sh.fill._xPr.find(qn("a:gradFill")).find(qn("a:gsLst"))
    mid = parse_xml(
        '<a:gs xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" pos="50000">'
        '<a:srgbClr val="3B82F6"/></a:gs>'
    )
    gsLst.insert(1, mid)
    _shadow(sh, blur=15, dist=0, alpha=0x29, color="2C1E74")
    sh.text_frame.word_wrap = True
    return sh


def pill(sl, x, y, w, h, text, *, fill=SURFACE, color=INK_DARK, size=10.5, weight=600):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.5
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, weight=weight, color=color, line=1.15, align=PP_ALIGN.CENTER, first=True)
    return sh


def circle(sl, x, y, d, text, *, fill=BLUE, color=WHITE, size=13, weight=700):
    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, text, size=size, weight=weight, color=color, line=1.0, align=PP_ALIGN.CENTER, first=True)
    return sh


def arrow(sl, x1, y1, x2, y2, *, color=BLUE_LIGHT, width=1.5, dashed=False):
    from pptx.oxml import parse_xml

    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    ln.append(parse_xml(
        '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'type="triangle" w="med" len="med"/>'
    ))
    if dashed:
        ln.append(parse_xml(
            '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'
        ))
    return cn


def line(sl, x1, y1, x2, y2, *, color=BORDER, width=1, dashed=False):
    from pptx.oxml import parse_xml

    cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    if dashed:
        cn.line._get_or_add_ln().append(parse_xml(
            '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'
        ))
    return cn


# ---------- 프레임 ----------
def frame(sl, *, chapter=None, page=None, source=None, dark=False):
    fg = WHITE if dark else MUTED
    if chapter:
        write(sl, M, HEADER_Y, 6.0, 0.3,
              [(chapter, dict(size=12, weight=600, color=(RGBColor(0xFF, 0xFF, 0xFF) if dark else MUTED), line=1.3))])
    img = LOGO_WHITE if dark else LOGO
    sl.shapes.add_picture(img, Inches(SLIDE_W - M - LOGO_W), Inches(0.44), Inches(LOGO_W), Inches(LOGO_H))
    if page is not None:
        write(sl, M, FOOT_Y, 2.0, 0.25,
              [(str(page), dict(size=10, weight=500, color=(WHITE if dark else MUTED), line=1.3))])
    if source:
        write(sl, SLIDE_W - M - 9.0, FOOT_Y, 9.0, 0.25,
              [(source, dict(size=9, weight=400, color=(WHITE if dark else MUTED), line=1.4))],
              align=PP_ALIGN.RIGHT)


def title_block(sl, headline, subtitle=None, *, size=32):
    write(sl, M, HEADLINE_Y, CW - 1.2, 0.63,
          [(headline, dict(size=size, weight=700, color=INK, line=1.2))])
    if subtitle:
        write(sl, M, SUBTITLE_Y, CW - 1.2, 0.4,
              [(subtitle, dict(size=16, weight=500, color=SUB, line=1.45))])


def content_slide(prs, chapter, headline, subtitle, page, source=None):
    sl = blank(prs)
    frame(sl, chapter=chapter, page=page, source=source)
    title_block(sl, headline, subtitle)
    return sl


def divider(prs, number, title, lead, page):
    sl = blank(prs, dark=True)
    frame(sl, chapter=number, page=page, dark=True)
    write(sl, M, 3.0, 11.0, 1.2, [(title, dict(size=44, weight=700, color=WHITE, line=1.15))])
    write(sl, M, 4.35, 11.0, 0.6, [(lead, dict(size=18, weight=500, color=RGBColor(0xC9, 0xD3, 0xE0), line=1.45))])
    return sl


# ---------- 차트 ----------
def style_chart(chart, *, colors=(BLUE,), label_fmt="0.0", labels=True, legend=False,
                axis_size=10, label_size=11, gap=60, y_max=None, y_min=None, number_format=None):
    chart.font.name = F400
    chart.font.size = Pt(axis_size)
    chart.font.color.rgb = SUB
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(10)
        chart.legend.font.name = F500
        chart.legend.font.color.rgb = SUB
    plot = chart.plots[0]
    plot.gap_width = gap
    for i, ser in enumerate(chart.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = colors[i % len(colors)]
        ser.format.line.fill.background()
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = label_fmt
        dl.number_format_is_linked = False
        dl.font.size = Pt(label_size)
        dl.font.name = F600
        dl.font.color.rgb = INK
        try:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception:
            pass
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = BORDER
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(axis_size)
    va.tick_labels.font.name = F400
    va.tick_labels.font.color.rgb = SUB
    if number_format:
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
    if y_max is not None:
        va.maximum_scale = y_max
    if y_min is not None:
        va.minimum_scale = y_min
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = BORDER
    ca.tick_labels.font.size = Pt(axis_size)
    ca.tick_labels.font.name = F500
    ca.tick_labels.font.color.rgb = SUB
    return chart


def add_chart(sl, kind, x, y, w, h, categories, series, **kw):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    gf = sl.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = gf.chart
    style_chart(chart, legend=len(series) > 1, **kw)
    return chart


COL = XL_CHART_TYPE.COLUMN_CLUSTERED
BAR = XL_CHART_TYPE.BAR_CLUSTERED
